#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pubmed_search.py — Recherche PubMed (E-utilities NCBI) et téléchargement des PDF
d'articles en libre accès.

Sources de PDF essayées, dans l'ordre :
  1. PMC Cloud Service (AWS S3, bucket public pmc-oa-opendata) — voie officielle
     depuis le retrait du service OA web/FTP de PMC (août 2026).
  2. URLs PDF libres listées par Europe PMC (API REST).
  3. Unpaywall (DOI) — versions OA chez l'éditeur, dépôts, preprints.
  4. OpenAlex (seulement si une clé API est fournie, obligatoire depuis 2026).
  5. Rendu PDF Europe PMC (souvent bloqué par Cloudflare, tenté une fois).
  6. Site de l'éditeur (DOI, liens LinkOut PubMed) : balise citation_pdf_url,
     motifs d'URL connus.

  Avec une clé Springer Nature Meta API : lien PDF officiel des articles Springer/BMC/
  Nature signalés en libre accès.

Passe JSON (--getjson), pour les articles restés sans PDF (le PDF reste prioritaire) :
  1. NCBI BioC API (PMC Open Access + manuscrits d'auteurs, texte intégral structuré).
  2. Europe PMC fullTextXML (JATS) converti en JSON.
  3. Springer Nature Open Access API (JATS, clé requise).
  4. API Elsevier (texte intégral JSON si l'article est libre pour la clé).
  5. Springer Nature Meta API (métadonnées + résumé seulement, en dernier recours).
  Tous les JSON ont le même schéma (voir jats_to_doc / bioc_to_doc).
  Karger : pas d'API publique de texte intégral (accès TDM sur contrat, par FTP).

Passe XML JATS (--getjats), pour tous les articles PMC : dossier JATS_<nom>/ « IA locale »
  (XML d'origine, Markdown avec tableaux complets, CSV par tableau, figures) depuis
  PMC S3, sinon Europe PMC (XML seul).

Mots-clés : <nom>_mots-cles.txt, une ligne par groupe de termes équivalents (MeSH + synonymes
officiels, mots-clés d'auteurs, correspondance PubMed, sigles des résumés) ; --no-keywords pour
ne pas le créer. --keywords-from-queries FICHIER : idem pour les termes d'une liste de requêtes.

Le script NE contourne PAS les protections anti-robot (CAPTCHA, preuve de travail
JavaScript, Cloudflare) : ces cas sont détectés et consignés dans error.txt.

Dépendances : pip install requests openpyxl

Exemples :
  python pubmed_search.py "diabetes[ti] AND review[pt]" --output Diabetes_review.xlsx \
         --max-results 50 --getfreepaper
  python pubmed_search.py "diabetes[ti] AND review[pt]" --output Diabetes_review.xlsx \
         --best-match --most-recent --getfreepaper --pass-number 2 --time-set 01-10-2026_20h45
"""

import argparse
import datetime as dt
import html as htmlmod
import json
import os
import re
import sys
import threading
import time
import unicodedata
import xml.etree.ElementTree as ET
from collections import Counter, OrderedDict
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import quote, urljoin, urlparse, unquote
import glob
import io
import shutil

try:
    import requests
except ImportError:
    sys.exit("Module 'requests' manquant : pip install requests")
try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
except ImportError:
    sys.exit("Module 'openpyxl' manquant : pip install openpyxl")
try:  # optionnel : lecture du texte des PDF (--import-pdf, contrôle API Elsevier)
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

VERSION = "1.10.1"
TOOL_NAME = "pubmed_search_py"
DEFAULT_EMAIL = "sergesawadogo@gmail.com"

EUTILS = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/"
IDCONV = "https://pmc.ncbi.nlm.nih.gov/tools/idconv/api/v1/articles/"
S3_BASE = "https://pmc-oa-opendata.s3.amazonaws.com/"
EPMC_SEARCH = "https://www.ebi.ac.uk/europepmc/webservices/rest/search"
UNPAYWALL = "https://api.unpaywall.org/v2/"
OPENALEX = "https://api.openalex.org/works"
HAL_API = "https://api.archives-ouvertes.fr/search/"
CORE_API = "https://api.core.ac.uk/v3/search/works"
ELSEVIER_ARTICLE = "https://api.elsevier.com/content/article/"
# préfixes DOI des revues hébergées sur ScienceDirect (Elsevier, Cell Press, JBC, JLR…)
ELSEVIER_DOI_PREFIXES = ("10.1016/", "10.1074/", "10.1194/", "10.3168/", "10.1053/", "10.1067/",
                         "10.1078/", "10.1006/", "10.1054/", "10.1157/", "10.1203/")
BIOC_PMC = "https://www.ncbi.nlm.nih.gov/research/bionlp/RESTful/pmcoa.cgi/BioC_json/{}/unicode"
EPMC_REST = "https://www.ebi.ac.uk/europepmc/webservices/rest/"
SPRINGER_META = "https://api.springernature.com/meta/v2/json"
SPRINGER_OA_JATS = "https://api.springernature.com/openaccess/jats"
# préfixes DOI Springer Nature (Springer, BMC, Nature, Palgrave, EPJ, Kluwer, Pleiades…)
SPRINGER_DOI_PREFIXES = ("10.1007/", "10.1186/", "10.1038/", "10.1057/", "10.1140/", "10.1245/",
                         "10.1023/", "10.1134/", "10.1365/", "10.1617/")
KARGER_DOI_PREFIX = "10.1159/"
# sources désactivables (--sources) ; les autres (PMC S3, Europe PMC, HAL, éditeur) restent actives
OPTIONAL_SOURCES = ("elsevier", "springer-oa", "springer-meta", "unpaywall", "core", "openalex")

BROWSER_HEADERS = {
    "User-Agent": ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/128.0 Safari/537.36"),
    "Accept": "text/html,application/xhtml+xml,application/pdf;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.8,fr;q=0.6",
}

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9\-]+(?:\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}")
ANTIBOT_MARKERS = (
    "just a moment", "cf-chl", "cf_chl", "challenge-platform", "captcha",
    "recaptcha", "hcaptcha", "proof of work", "proof-of-work", "anubis",
    "preparing to download", "are you a robot", "verify you are human",
    "access denied", "request unsuccessful", "incapsula", "perimeterx",
    "ddos-guard", "bot detection",
)
PMC_WEB_RE = re.compile(r"(pmc\.ncbi\.nlm\.nih\.gov|ncbi\.nlm\.nih\.gov/pmc)", re.I)
MAX_PDF_BYTES = 150 * 1024 * 1024

PRINT_LOCK = threading.Lock()
STOP = threading.Event()


def log(msg=""):
    with PRINT_LOCK:
        print(msg, flush=True)


def short_err(e):
    """Message d'erreur court (sans l'URL, souvent longue de plusieurs milliers de caractères)."""
    resp = getattr(e, "response", None)
    if resp is not None and getattr(resp, "status_code", None):
        names = {429: "trop de requêtes", 500: "erreur interne du serveur", 502: "passerelle",
                 503: "service indisponible", 504: "délai dépassé côté serveur"}
        return f"HTTP {resp.status_code} ({names.get(resp.status_code, resp.reason or 'erreur')})"
    return f"{type(e).__name__} : {str(e)[:120]}"


# services mis en pause après plusieurs échecs serveur consécutifs (les autres sources continuent)
BREAKER_SERVICES = {"europepmc": "Europe PMC", "openalex": "OpenAlex", "unpaywall": "Unpaywall",
                    "hal": "HAL", "crossref": "Crossref", "bioc": "NCBI BioC", "idconv": "ID converter PMC",
                    "meshapi": "API MeSH (NLM)"}
BREAKER_FAILS = 3          # échecs consécutifs (après les nouvelles tentatives) avant la pause
BREAKER_PAUSE = (120, 300, 600)   # durée des pauses successives (s)


def chunks(seq, n):
    seq = list(seq)
    for i in range(0, len(seq), n):
        yield seq[i:i + n]


def clean_xl(v):
    if isinstance(v, str):
        return ILLEGAL_CHARACTERS_RE.sub("", v)
    return v


def ascii_slug(s, keep=r"A-Za-z0-9\-"):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode()
    s = re.sub(r"\s+", "-", s.strip())
    return re.sub(r"[^%s]" % keep, "", s)


def fmt_duration(sec):
    sec = int(round(sec))
    h, r = divmod(sec, 3600)
    m, s = divmod(r, 60)
    return f"{h}h{m:02d}m{s:02d}s" if h else f"{m}m{s:02d}s"


# --------------------------------------------------------------------------- #
# Client HTTP : compteur de requêtes, limitation de débit, reprises
# --------------------------------------------------------------------------- #
class Http:
    def __init__(self, email, api_key=None, openalex_key=None, timeout=40, elsevier_key=None,
                 core_key=None, springer_oa_key=None, springer_meta_key=None, sources=None):
        self.email = email
        self.api_key = api_key
        self.sources = set(OPTIONAL_SOURCES if sources is None else sources)
        # une source décochée = clé ignorée (aucune requête vers ce service)
        self.openalex_key = openalex_key if "openalex" in self.sources else None
        self.elsevier_key = elsevier_key if "elsevier" in self.sources else None
        self.core_key = core_key if "core" in self.sources else None
        self.springer_oa_key = springer_oa_key if "springer-oa" in self.sources else None
        self.springer_meta_key = springer_meta_key if "springer-meta" in self.sources else None
        self.elsevier_insttoken = None
        self.els_denied = False   # clé Elsevier valide mais sans droits (AUTHENTICATION_ERROR)
        self.timeout = timeout
        self.intervals = {
            "ncbi": 0.11 if api_key else 0.35,   # 10 req/s avec clé, 3 sans
            "idconv": 0.35, "europepmc": 0.12, "unpaywall": 0.12,
            "openalex": 0.12, "s3": 0.0, "publisher": 1.0, "handle": 0.2, "elsevier": 0.15, "hal": 0.3, "core": 6.5, "meshapi": 0.2, "core_dl": 1.0,
            "repository": 1.0,
            "bioc": 0.34,          # NCBI BioC (pas de limite publiée : 3 req/s par prudence)
            "springer": 1.0,       # forfait gratuit Springer Nature : quota journalier limité
            "crossref": 0.2,       # API publique Crossref (« polite pool » avec mailto)
        }

        self.counts = Counter()
        self.skipped = Counter()      # requêtes évitées pendant la pause d'un service
        self.paused_log = []          # (service, durée) : pour les statistiques et error.txt
        self._fails = Counter()
        self._down_until = {}
        self._pauses = Counter()
        self._lock = threading.Lock()
        self._next = {}
        self._klocks = {}
        self._local = threading.local()

    def use(self, source):
        return source in self.sources

    def elsevier_headers(self):
        h = {"X-ELS-APIKey": self.elsevier_key or ""}
        if getattr(self, "elsevier_insttoken", None):
            h["X-ELS-Insttoken"] = self.elsevier_insttoken
        return h

    def _session(self):
        s = getattr(self._local, "s", None)
        if s is None:
            s = requests.Session()
            s.headers["User-Agent"] = f"{TOOL_NAME}/{VERSION} (mailto:{self.email})"
            self._local.s = s
        return s

    def _throttle(self, service, url):
        interval = self.intervals.get(service, 0.5)
        if interval <= 0:
            return
        key = urlparse(url).netloc if service in ("publisher", "repository") else service
        with self._lock:
            kl = self._klocks.setdefault(key, threading.Lock())
        with kl:
            now = time.monotonic()
            nxt = self._next.get(key, 0.0)
            if nxt > now:
                time.sleep(nxt - now)
            self._next[key] = time.monotonic() + interval

    @staticmethod
    def _unavailable(url, why):
        """Réponse 503 locale (aucune requête envoyée) pendant la pause d'un service."""
        r = requests.Response()
        r.status_code = 503
        r.reason = why
        r.url = url
        r._content = b""
        r.headers["X-PubMed-Search"] = "service en pause"
        return r

    def is_paused(self, service):
        return self._down_until.get(service, 0) > time.monotonic()

    def _breaker(self, service, failed, detail=""):
        if service not in BREAKER_SERVICES:
            return
        with self._lock:
            if not failed:
                self._fails[service] = 0
                return
            self._fails[service] += 1
            if self._fails[service] < BREAKER_FAILS or self.is_paused(service):
                return
            k = self._pauses[service]
            pause = BREAKER_PAUSE[min(k, len(BREAKER_PAUSE) - 1)]
            self._pauses[service] += 1
            self._down_until[service] = time.monotonic() + pause
            self._fails[service] = BREAKER_FAILS - 1   # après la pause : un seul essai de reprise
            self.paused_log.append((service, pause))
        log(f"  ⏸ {BREAKER_SERVICES[service]} ne répond pas ({detail or 'erreur serveur'}, "
            f"{BREAKER_FAILS} échecs de suite) : service mis en pause {pause // 60} min, "
            "les autres sources continuent.")

    def request(self, method, url, service, retries=3, **kw):
        kw.setdefault("timeout", self.timeout)
        if self.is_paused(service):
            with self._lock:
                self.skipped[service] += 1
            return self._unavailable(url, "service en pause après des erreurs répétées")
        if self._fails.get(service, 0) >= BREAKER_FAILS - 1 and service in BREAKER_SERVICES:
            retries = 0   # essai de reprise : pas d'attente supplémentaire s'il échoue encore
        for attempt in range(retries + 1):
            self._throttle(service, url)
            with self._lock:
                self.counts[service] += 1
            try:
                r = self._session().request(method, url, **kw)
            except requests.RequestException as e:
                if attempt < retries:
                    time.sleep(2 ** attempt)
                    continue
                self._breaker(service, True, type(e).__name__)
                raise
            if r.status_code in (429, 500, 502, 503, 504) and attempt < retries:
                ra = r.headers.get("Retry-After", "")
                wait = float(ra) if ra.isdigit() else 2 ** (attempt + 1)
                r.close()
                time.sleep(min(wait, 60))
                continue
            self._breaker(service, r.status_code in (500, 502, 503, 504), f"HTTP {r.status_code}")
            return r

    def ncbi(self, **params):
        p = {"tool": TOOL_NAME, "email": self.email}
        if self.api_key:
            p["api_key"] = self.api_key
        p.update(params)
        return p

    @property
    def total(self):
        return sum(self.counts.values())


# --------------------------------------------------------------------------- #
# Modèle d'article
# --------------------------------------------------------------------------- #
class Article:
    def __init__(self, pmid):
        self.pmid = pmid
        self.doi = ""
        self.pmcid = ""
        self.title = ""
        self.journal = ""
        self.journal_abbr = ""
        self.volume = ""
        self.issue = ""
        self.month = ""
        self.pages = ""
        self.year = ""
        self.pub_types = []
        self.authors = []          # [{"name","last","initials","emails"}]
        self.free = "Non"          # Free PMC article / Free article / PMC (embargo ?) / Non
        self.rank_best = None
        self.rank_recent = None
        self.pdf_rel = ""
        self.pdf_source = ""
        self.pdf_pass = None
        self.json_rel = ""
        self.json_source = ""
        self.json_level = ""       # "texte intégral" / "résumé"
        self.json_attempts = []    # (source, url, raison)
        self.abstract = ""
        self.mesh = []             # [(UI, descripteur)]
        self.keywords = []         # mots-clés des auteurs
        self.jats_rel = ""         # Markdown de l'article dans JATS_<nom>/
        self.jats_source = ""
        self.jats_info = ""
        self.jats_note = ""
        self.springer_meta = None  # notice Springer Meta API (réutilisée par la passe JSON)
        self.oa_hint = False       # une source externe signale une version libre
        self.epmc_pdf_urls = []
        self.landing_urls = []
        self.openalex_pdf_urls = []
        self.attempts = []         # (passe, source, url, raison)
        self._lock = threading.Lock()

    def note(self, pass_no, source, url, reason):
        with self._lock:
            self.attempts.append((pass_no, source, url or "", reason))

    def jnote(self, source, url, reason):
        with self._lock:
            self.json_attempts.append((source, url or "", reason))

    @property
    def is_springer(self):
        return bool(self.doi) and self.doi.lower().startswith(SPRINGER_DOI_PREFIXES)

    @property
    def is_elsevier(self):
        return bool(self.doi) and self.doi.lower().startswith(ELSEVIER_DOI_PREFIXES)

    @property
    def free_like(self):
        return self.free != "Non" or self.oa_hint

    def author_display(self, a):
        return a["name"]

    @property
    def first_author(self):
        return self.authors[0]["name"] if self.authors else ""

    @property
    def last_author(self):
        return self.authors[-1]["name"] if len(self.authors) > 1 else ""

    @property
    def other_authors(self):
        return "; ".join(a["name"] for a in self.authors[1:-1]) if len(self.authors) > 2 else ""

    @property
    def emails(self):
        out = []
        for a in self.authors:
            if a["emails"]:
                out.append(f"{a['name']}: {', '.join(a['emails'])}")
        return "; ".join(out)

    def pdf_filename(self):
        if self.authors:
            a = self.authors[0]
            last = ascii_slug(a["last"]) or "Anonyme"
            ini = ascii_slug(a["initials"], keep=r"A-Za-z")
            who = f"{last}_{ini}" if ini else last
        else:
            who = "Anonyme"
        year = self.year or "nd"
        return f"{self.pmid}-{who}({year}).pdf"

    @property
    def sort_label(self):
        if self.rank_best and self.rank_recent:
            return f"Les deux (BM #{self.rank_best}, MR #{self.rank_recent})"
        if self.rank_best:
            return f"Best match #{self.rank_best}"
        if self.rank_recent:
            return f"Most recent #{self.rank_recent}"
        return ""


# --------------------------------------------------------------------------- #
# PubMed : ESearch / EFetch / ELink
# --------------------------------------------------------------------------- #
# erreurs passagères du moteur de recherche NCBI (renvoyées avec un HTTP 200)
NCBI_TRANSIENT = ("backend failed", "database is not supported", "temporarily", "unavailable",
                  "timeout", "try again", "internal error", "cannot process")


def esearch(http, term, sort, retmax):
    data = http.ncbi(db="pubmed", term=term, sort=sort, retmax=retmax, retmode="json")
    waits = (10, 30, 60, 120)
    for attempt in range(len(waits) + 1):
        r = http.request("POST", EUTILS + "esearch.fcgi", "ncbi", data=data)
        r.raise_for_status()
        try:
            j = r.json().get("esearchresult", {})
        except ValueError:
            j = {"ERROR": "réponse illisible (backend NCBI)"}
        err = str(j.get("ERROR", ""))
        if err and any(k in err.lower() for k in NCBI_TRANSIENT) and attempt < len(waits) and not STOP.is_set():
            log(f"  ⏳ PubMed ne répond pas correctement ({err[:90]}) : panne passagère des serveurs NCBI, "
                f"nouvel essai dans {waits[attempt]} s ({attempt + 1}/{len(waits)})…")
            time.sleep(waits[attempt])
            continue
        break
    if "ERROR" in j:
        if any(k in str(j["ERROR"]).lower() for k in NCBI_TRANSIENT):
            raise RuntimeError(f"ESearch : {j['ERROR']} — panne des serveurs NCBI qui dure depuis près de "
                               "4 min : relancez la recherche plus tard. Si l'erreur persiste, testez la même "
                               "requête sur pubmed.ncbi.nlm.nih.gov.")
        raise RuntimeError(f"ESearch : {j['ERROR']}")
    warn = j.get("warninglist") or {}
    for k in ("phrasesignored", "quotedphrasesnotfound", "outputmessages"):
        if warn.get(k):
            log(f"  ⚠ PubMed ({k}) : {warn[k]}")
    errl = j.get("errorlist") or {}
    if errl.get("phrasesnotfound"):
        log(f"  ⚠ Termes non trouvés : {errl['phrasesnotfound']}")
    return int(j.get("count", 0)), j.get("idlist", []), j.get("querytranslation", "")


def _txt(el):
    return re.sub(r"\s+", " ", "".join(el.itertext())).strip() if el is not None else ""


def _year_of(node):
    if node is None:
        return ""
    for path in (".//PubDate/Year", ".//ArticleDate/Year"):
        y = node.findtext(path)
        if y and y.strip().isdigit():
            return y.strip()
    md = node.findtext(".//PubDate/MedlineDate") or ""
    m = re.search(r"(19|20)\d{2}", md)
    if m:
        return m.group(0)
    y = node.findtext(".//PubMedPubDate[@PubStatus='pubmed']/Year")
    return (y or "").strip()


def _parse_authors(author_els):
    out = []
    for au in author_els:
        last = (au.findtext("LastName") or "").strip()
        ini = (au.findtext("Initials") or "").strip()
        fore = (au.findtext("ForeName") or "").strip()
        coll = _txt(au.find("CollectiveName"))
        if not ini and fore:
            ini = "".join(p[0] for p in re.split(r"[\s\-]+", fore) if p)
        name = f"{last} {ini}".strip() if last else coll
        if not name:
            continue
        affs = " ".join(_txt(a) for a in au.findall("AffiliationInfo/Affiliation"))
        emails = []
        for e in EMAIL_RE.findall(affs):
            e = e.rstrip(".")
            if e.lower() not in [x.lower() for x in emails]:
                emails.append(e)
        out.append({"name": name, "last": last or coll, "initials": ini if last else "",
                    "emails": emails})
    return out


def parse_pubmed_article(el):
    mc = el.find("MedlineCitation")
    a = Article((mc.findtext("PMID") or "").strip())
    art = mc.find("Article")
    a.title = _txt(art.find("ArticleTitle")) or _txt(art.find("VernacularTitle"))
    a.journal = _txt(art.find("Journal/Title"))
    # éléments de la référence bibliographique (style Vancouver / NLM)
    a.journal_abbr = (mc.findtext("MedlineJournalInfo/MedlineTA") or
                      art.findtext("Journal/ISOAbbreviation") or "").strip()
    ji = art.find("Journal/JournalIssue")
    a.volume = (ji.findtext("Volume") or "").strip() if ji is not None else ""
    a.issue = (ji.findtext("Issue") or "").strip() if ji is not None else ""
    a.month = (ji.findtext("PubDate/Month") or "").strip() if ji is not None else ""
    a.pages = (art.findtext("Pagination/MedlinePgn") or "").strip()
    if not a.pages:
        for eloc in art.findall("ELocationID"):
            if (eloc.get("EIdType") or "").lower() == "pii":
                a.pages = (eloc.text or "").strip()
                break
    a.year = _year_of(art.find("Journal/JournalIssue")) or _year_of(art) or _year_of(el)
    a.pub_types = [_txt(p) for p in art.findall("PublicationTypeList/PublicationType")]
    a.authors = _parse_authors(art.findall("AuthorList/Author"))
    ab = art.find("Abstract")
    if ab is not None:
        a.abstract = " ".join(_txt(t) for t in ab.findall("AbstractText"))
    for mh in mc.findall("MeshHeadingList/MeshHeading"):
        d = mh.find("DescriptorName")
        if d is not None and _txt(d):
            a.mesh.append((d.get("UI") or "", _txt(d)))
    for kw in mc.findall("KeywordList/Keyword"):
        t = _txt(kw)
        if t and t not in a.keywords:
            a.keywords.append(t)
    idlist = el.find("PubmedData/ArticleIdList")
    if idlist is not None:
        for aid in idlist.findall("ArticleId"):
            t = (aid.get("IdType") or "").lower()
            v = (aid.text or "").strip()
            if t == "doi" and not a.doi:
                a.doi = v
            elif t == "pmc" and not a.pmcid:
                a.pmcid = v if v.upper().startswith("PMC") else "PMC" + v
    if not a.doi:
        for eloc in art.findall("ELocationID"):
            if (eloc.get("EIdType") or "").lower() == "doi":
                a.doi = (eloc.text or "").strip()
                break
    return a


def parse_book_article(el):
    bd = el.find("BookDocument")
    a = Article((bd.findtext("PMID") or "").strip())
    a.title = _txt(bd.find("ArticleTitle")) or _txt(bd.find("Book/BookTitle"))
    a.journal = _txt(bd.find("Book/Publisher/PublisherName"))
    a.year = _year_of(bd.find("Book")) or ""
    m = re.search(r"(19|20)\d{2}", _txt(bd.find("Book/PubDate")))
    if not a.year and m:
        a.year = m.group(0)
    a.pub_types = [_txt(p) for p in bd.findall("PublicationType")] or ["Book chapter"]
    a.authors = _parse_authors(bd.findall("AuthorList/Author"))
    for aid in el.findall("PubmedBookData/ArticleIdList/ArticleId"):
        t = (aid.get("IdType") or "").lower()
        if t == "doi" and not a.doi:
            a.doi = (aid.text or "").strip()
        elif t == "pmc" and not a.pmcid:
            a.pmcid = (aid.text or "").strip()
    return a


def efetch_articles(http, pmids, quiet=False):
    out = {}
    for batch in chunks(pmids, 200):
        r = http.request("POST", EUTILS + "efetch.fcgi", "ncbi",
                         data=http.ncbi(db="pubmed", id=",".join(batch), retmode="xml"))
        r.raise_for_status()
        root = ET.fromstring(r.content)
        for el in root.findall("PubmedArticle"):
            try:
                a = parse_pubmed_article(el)
                out[a.pmid] = a
            except Exception as e:  # noqa
                log(f"  ⚠ Erreur d'analyse XML : {e}")
        for el in root.findall("PubmedBookArticle"):
            try:
                a = parse_book_article(el)
                out[a.pmid] = a
            except Exception as e:  # noqa
                log(f"  ⚠ Erreur d'analyse XML (livre) : {e}")
        if not quiet:
            log(f"  métadonnées : {len(out)}/{len(pmids)}")
    return out


def free_fulltext_set(http, pmids):
    """PMIDs portant le filtre PubMed 'Free full text' (= mention Free article / Free PMC article)."""
    res = set()
    for batch in chunks(pmids, 250):
        term = "(" + " OR ".join(f"{p}[pmid]" for p in batch) + ") AND free full text[sb]"
        _, ids, _ = esearch(http, term, "relevance", len(batch))
        res.update(ids)
    return res


def elink_prlinks(http, articles):
    """Liens 'Full text' (LinkOut) que PubMed affiche vers l'éditeur."""
    for batch in chunks(articles, 100):
        try:
            r = http.request("POST", EUTILS + "elink.fcgi", "ncbi",
                             data=http.ncbi(dbfrom="pubmed", cmd="prlinks", retmode="json",
                                            id=[a.pmid for a in batch]))
            r.raise_for_status()
            j = r.json()
        except Exception as e:  # noqa
            log(f"  ⚠ ELink prlinks : {short_err(e)}")
            continue
        byid = {a.pmid: a for a in batch}
        for ls in j.get("linksets", []):
            for item in ls.get("idurllist", []):
                a = byid.get(str(item.get("id")))
                if not a:
                    continue
                for o in item.get("objurls", []):
                    url = (o.get("url") or {}).get("value", "")
                    if url and url not in a.landing_urls:
                        a.landing_urls.append(url)


# --------------------------------------------------------------------------- #
# Enrichissement : ID converter, Europe PMC, OpenAlex
# --------------------------------------------------------------------------- #
def elink_pmc_fill(http, articles):
    """PMCID par E-utilities (ELink pubmed → pmc) : repli quand l'ID converter est refusé."""
    todo = [a for a in articles if not a.pmcid]
    found = 0
    for batch in chunks(todo, 200):
        try:
            r = http.request("POST", EUTILS + "elink.fcgi", "ncbi",
                             data=http.ncbi(dbfrom="pubmed", db="pmc", linkname="pubmed_pmc",
                                            retmode="json", id=[a.pmid for a in batch]))
            r.raise_for_status()
            j = r.json()
        except Exception as e:  # noqa
            log(f"  ⚠ ELink pubmed→pmc : {short_err(e)}")
            continue
        byid = {a.pmid: a for a in batch}
        for ls in j.get("linksets", []):
            ids = [str(i) for i in ls.get("ids", [])]
            if len(ids) != 1 or ids[0] not in byid:
                continue
            for db in ls.get("linksetdbs", []):
                if db.get("linkname") == "pubmed_pmc" and db.get("links"):
                    byid[ids[0]].pmcid = f"PMC{db['links'][0]}"
                    found += 1
                    break
    return found


def idconv_fill(http, articles):
    todo = [a for a in articles if not a.pmcid or not a.doi]
    for batch in chunks(todo, 200):
        try:
            r = http.request("GET", IDCONV, "idconv",
                             params={"ids": ",".join(a.pmid for a in batch), "idtype": "pmid",
                                     "format": "json", "tool": TOOL_NAME, "email": http.email})
            r.raise_for_status()
            recs = r.json().get("records", [])
        except Exception as e:  # noqa
            code = getattr(getattr(e, "response", None), "status_code", None)
            if code in (403, 429):
                # Refus du site PMC (protection anti-robots) : on passe par E-utilities, qui
                # donne les mêmes PMCID ; les DOI viennent d'EFetch, Europe PMC et OpenAlex.
                n = elink_pmc_fill(http, todo)
                log(f"  ℹ ID converter PMC indisponible (HTTP {code}) : PMCID recherchés via "
                    f"E-utilities à la place ({n} trouvé(s)).")
                return
            log(f"  ⚠ ID converter PMC : {short_err(e)}")
            continue
        byid = {a.pmid: a for a in batch}
        for rec in recs:
            a = byid.get(str(rec.get("pmid") or rec.get("requested-id")))
            if not a:
                continue
            if rec.get("pmcid") and not a.pmcid:
                a.pmcid = rec["pmcid"]
            if rec.get("doi") and not a.doi:
                a.doi = rec["doi"]


def europepmc_enrich(http, articles):
    batches = list(chunks(articles, 100))
    done, failed = 0, 0
    for i, batch in enumerate(batches, 1):
        if http.is_paused("europepmc"):
            n_left = sum(len(b) for b in batches[i - 1:])
            log(f"  ℹ Europe PMC en pause : {n_left} article(s) non enrichis par Europe PMC "
                "(les PDF restent cherchés sur PMC S3, Unpaywall, OpenAlex, HAL, éditeurs).")
            return
        q = "SRC:MED AND (" + " OR ".join(f"EXT_ID:{a.pmid}" for a in batch) + ")"
        try:
            r = http.request("GET", EPMC_SEARCH, "europepmc", retries=1,
                             params={"query": q, "resultType": "core", "format": "json",
                                     "pageSize": 1000})
            r.raise_for_status()
            results = r.json().get("resultList", {}).get("result", [])
        except Exception as e:  # noqa
            failed += 1
            log(f"  ⚠ Europe PMC : lot {i}/{len(batches)} non obtenu — {short_err(e)}")
            continue
        done += len(batch)
        log(f"  Europe PMC : {done}/{len(articles)}")
        byid = {a.pmid: a for a in batch}
        for res in results:
            a = byid.get(str(res.get("pmid") or res.get("id")))
            if not a:
                continue
            if res.get("pmcid") and not a.pmcid:
                a.pmcid = res["pmcid"]
            if res.get("doi") and not a.doi:
                a.doi = res["doi"]
            if res.get("isOpenAccess") == "Y":
                a.oa_hint = True
            for ft in (res.get("fullTextUrlList") or {}).get("fullTextUrl", []):
                avail = (ft.get("availabilityCode") or "").upper()
                url = ft.get("url", "")
                if avail not in ("OA", "F") or not url:
                    continue
                a.oa_hint = True
                if (ft.get("documentStyle") or "").lower() == "pdf":
                    if url not in a.epmc_pdf_urls:
                        a.epmc_pdf_urls.append(url)
                elif url not in a.landing_urls:
                    a.landing_urls.append(url)


def openalex_enrich(http, articles):
    if not http.openalex_key:
        return
    for k, batch in enumerate(chunks(articles, 50), 1):
        if k > 1 and (k - 1) % 4 == 0:
            log(f"  OpenAlex : {min((k - 1) * 50, len(articles))}/{len(articles)}")
        try:
            r = http.request("GET", OPENALEX, "openalex",
                             params={"filter": "pmid:" + "|".join(a.pmid for a in batch),
                                     "per-page": 50, "api_key": http.openalex_key,
                                     "select": "ids,open_access,best_oa_location,locations"})
            r.raise_for_status()
            results = r.json().get("results", [])
        except Exception as e:  # noqa
            log(f"  ⚠ OpenAlex : {short_err(e)}")
            if http.is_paused("openalex"):
                log("  ℹ OpenAlex en pause : enrichissement OpenAlex arrêté pour cette recherche.")
                return
            continue
        byid = {a.pmid: a for a in batch}
        for w in results:
            pm = str((w.get("ids") or {}).get("pmid", "")).rstrip("/").split("/")[-1]
            a = byid.get(pm)
            if not a:
                continue
            if (w.get("open_access") or {}).get("is_oa"):
                a.oa_hint = True
            locs = [w.get("best_oa_location")] + (w.get("locations") or [])
            for loc in locs:
                if loc and loc.get("is_oa") and loc.get("pdf_url"):
                    if loc["pdf_url"] not in a.openalex_pdf_urls:
                        a.openalex_pdf_urls.append(loc["pdf_url"])


# --------------------------------------------------------------------------- #
# Téléchargement
# --------------------------------------------------------------------------- #
def looks_antibot(text):
    t = text[:20000].lower()
    return any(m in t for m in ANTIBOT_MARKERS)


def find_pdf_links(page, base_url, doi=""):
    links = []

    def add(u):
        if not u:
            return
        u = urljoin(base_url, htmlmod.unescape(u.strip()))
        if u.startswith("http") and u not in links:
            links.append(u)

    for pat in (r'<meta[^>]+name=["\']citation_pdf_url["\'][^>]*content=["\']([^"\']+)',
                r'<meta[^>]+content=["\']([^"\']+)["\'][^>]*name=["\']citation_pdf_url["\']',
                r'<link[^>]+type=["\']application/pdf["\'][^>]*href=["\']([^"\']+)',
                r'<link[^>]+href=["\']([^"\']+)["\'][^>]*type=["\']application/pdf["\']'):
        for m in re.finditer(pat, page, re.I):
            add(m.group(1))
    m = re.search(r'<meta[^>]+http-equiv=["\']refresh["\'][^>]*content=["\'][^"\']*url=([^"\'>]+)',
                  page, re.I)
    if m:
        add(m.group(1).strip("'\" "))
    m = re.search(r'name=["\']redirectURL["\'][^>]*value=["\']([^"\']+)', page, re.I)
    if m:  # Elsevier linkinghub
        add(unquote(m.group(1)))
    for u in publisher_guesses(base_url, doi):
        add(u)
    for m in re.finditer(r'href=["\']([^"\']+?\.pdf(?:\?[^"\']*)?)["\']', page, re.I):
        if len(links) >= 8:
            break
        add(m.group(1))
    return links


def url_key(u):
    """Clé de dédoublonnage : sans schéma ni www ; DOI insensible à la casse."""
    k = re.sub(r"^https?://(www\.)?", "", u.strip())
    return k.lower() if k.lower().startswith(("doi.org/", "dx.doi.org/")) else k


def publisher_guesses(url, doi):
    p = urlparse(url)
    host, path = p.netloc.lower(), p.path
    clean = f"{p.scheme}://{p.netloc}{path}".rstrip("/")
    g = []
    if "mdpi.com" in host and not path.endswith("/pdf"):
        g.append(clean + "/pdf")
    if "frontiersin.org" in host and doi:
        g.append(f"https://www.frontiersin.org/articles/{doi}/pdf")
    if "plos.org" in host and doi:
        seg = path.strip("/").split("/")[0] if path.strip("/") else "plosone"
        g.append(f"https://journals.plos.org/{seg}/article/file?id={doi}&type=printable")
    if any(h in host for h in ("springer", "biomedcentral", "springeropen")) and doi:
        g.append(f"https://link.springer.com/content/pdf/{doi}.pdf")
    if "nature.com" in host and "/articles/" in path and not path.endswith(".pdf"):
        g.append(clean + ".pdf")
    if "wiley.com" in host and doi:
        g.append(f"https://{host}/doi/pdfdirect/{doi}")
    if any(h in host for h in ("tandfonline.com", "sagepub.com", "pubs.acs.org", "ahajournals.org",
                               "atsjournals.org", "nejm.org", "science.org", "liebertpub.com",
                               "physiology.org", "annualreviews.org")) and doi:
        g.append(f"https://{host}/doi/pdf/{doi}")
    if ("bmj.com" in host or "biorxiv.org" in host or "medrxiv.org" in host) \
            and not path.endswith(".pdf"):
        g.append(re.sub(r"\.full$", "", clean) + ".full.pdf")
    pii = extract_pii(url)
    if pii and any(h in host for h in ("sciencedirect", "elsevier", "cell.com", "jbc.org")):
        if "cell.com" in host or "linkinghub" in host:
            g.append("https://www.cell.com/action/showPdf?pii=" + quote(format_pii(pii), safe=""))
        g.append(f"https://www.sciencedirect.com/science/article/pii/{pii}/pdfft?isDTMRedir=true&download=true")
    return g


def extract_pii(url):
    """PII Elsevier compact (S0021925820654576) depuis /pii/… ou ?pii=…"""
    u = unquote(url)
    m = re.search(r"(?:/pii/|[?&]pii=)([^/?&#]+)", u, re.I)
    if not m:
        return ""
    pii = re.sub(r"[^0-9A-Za-z]", "", m.group(1)).upper()
    return pii if re.fullmatch(r"[SB][0-9X]{16}", pii) else ""


def format_pii(pii):
    """S0960982206011079 -> S0960-9822(06)01107-9 (forme attendue par cell.com)."""
    if len(pii) == 17 and pii[0] == "S":
        return f"S{pii[1:5]}-{pii[5:9]}({pii[9:11]}){pii[11:16]}-{pii[16]}"
    return pii


class Downloader:
    def __init__(self, http, pdf_dir, pdf_dirname):
        self.http = http
        self.pdf_dir = pdf_dir
        self.pdf_dirname = pdf_dirname
        self.source_ok = Counter()

    # -- requête unitaire --------------------------------------------------- #
    def try_url(self, a, pass_no, source, url, service, follow_html=True, referer=None,
                extra_headers=None):
        """Retourne (ok, html, final_url)."""
        if STOP.is_set():
            return False, None, None
        key = url_key(url)
        seen = getattr(a, "_tried", None)
        if seen is not None:
            if key in seen:
                return False, None, None
            seen.add(key)
        if urlparse(url).netloc.lower() in ("hdl.handle.net", "handle.net"):
            real = self.resolve_handle(a, pass_no, source, url)
            if not real:
                return False, None, None
            url = real
            if seen is not None:
                seen.add(url_key(url))
        if PMC_WEB_RE.search(url) and service != "s3":
            a.note(pass_no, source, url, "ignoré : site web PMC protégé par une preuve de "
                                         "travail JavaScript (PDF cherché via PMC S3)")
            return False, None, None
        headers = dict(BROWSER_HEADERS) if service == "publisher" else {}
        headers.update(extra_headers or {})
        if referer:
            headers["Referer"] = referer
        try:
            r = self.http.request("GET", url, service, retries=(1 if service in ("publisher", "repository") else 2),
                                  stream=True, headers=headers, allow_redirects=True)
        except requests.Timeout:
            a.note(pass_no, source, url, "délai dépassé (timeout)")
            return False, None, None
        except requests.RequestException as e:
            a.note(pass_no, source, url, f"erreur réseau : {type(e).__name__}: {str(e)[:150]}")
            return False, None, None
        with r:
            ctype = r.headers.get("Content-Type", "")
            if r.status_code >= 400:
                body = ""
                try:
                    body = r.raw.read(30000, decode_content=True).decode("utf-8", "replace")
                except Exception:  # noqa
                    pass
                extra = " — page anti-robot (Cloudflare/CAPTCHA)" if looks_antibot(body) else ""
                if not extra and service not in ("publisher", "s3") and body.strip():
                    msg = re.sub(r"<[^>]+>", " ", body)
                    msg = re.sub(r"\s+", " ", msg).strip()[:220]
                    extra = f" — réponse : {msg}"
                a.note(pass_no, source, url, f"HTTP {r.status_code}{extra}")
                return False, None, None
            data = bytearray()
            try:
                it = r.iter_content(65536)
                for chunk in it:
                    data.extend(chunk)
                    if len(data) >= 2048:
                        break
                is_pdf = b"%PDF" in bytes(data[:1024])
                limit = MAX_PDF_BYTES if is_pdf else 3 * 1024 * 1024
                for chunk in it:
                    data.extend(chunk)
                    if len(data) > limit:
                        break
            except requests.RequestException as e:
                a.note(pass_no, source, url, f"transfert interrompu : {type(e).__name__}")
                return False, None, None
            if is_pdf:
                if len(data) > MAX_PDF_BYTES:
                    a.note(pass_no, source, url, "PDF trop volumineux (>150 Mo)")
                    return False, None, None
                if len(data) < 4000:
                    a.note(pass_no, source, url, f"PDF suspect (trop petit : {len(data)} octets)")
                    return False, None, None
                if service == "elsevier" and PdfReader is not None:
                    try:
                        npages = len(PdfReader(io.BytesIO(bytes(data))).pages)
                    except Exception:  # noqa
                        npages = 0
                    if npages == 1:
                        a.note(pass_no, source, url, "API Elsevier : 1re page seulement "
                                                     "(article non libre pour cette clé)")
                        return False, None, None
                fname = a.pdf_filename()
                tmp = os.path.join(self.pdf_dir, fname + ".part")
                with open(tmp, "wb") as f:
                    f.write(data)
                os.replace(tmp, os.path.join(self.pdf_dir, fname))
                a.pdf_rel = f"{self.pdf_dirname}/{fname}"
                a.pdf_source = source
                a.pdf_pass = pass_no
                return True, None, None
            text = bytes(data).decode(r.encoding or "utf-8", "replace")
            final = r.url
        if looks_antibot(text):
            a.note(pass_no, source, url, "page anti-robot reçue au lieu du PDF "
                                         "(JavaScript/CAPTCHA/Cloudflare) — non contournée")
            return False, None, None
        a.note(pass_no, source, url, f"HTML reçu au lieu d'un PDF ({ctype.split(';')[0] or '?'})")
        if follow_html:
            for link in find_pdf_links(text, final, a.doi)[:4]:
                if link == url:
                    continue
                ok, _, _ = self.try_url(a, pass_no, source + " (lien trouvé dans la page)", link,
                                        service if service != "s3" else "publisher",
                                        follow_html=False, referer=final)
                if ok:
                    return True, None, None
        return False, text, final

    def resolve_handle(self, a, pass_no, source, url):
        """hdl.handle.net/2066/127657 -> URL réelle du dépôt via l'API Handle."""
        hid = urlparse(url).path.strip("/")
        api = f"https://hdl.handle.net/api/handles/{hid}"
        try:
            r = self.http.request("GET", api, "handle", retries=1)
            vals = r.json().get("values", []) if r.status_code == 200 else []
        except (requests.RequestException, ValueError) as e:
            a.note(pass_no, source, url, f"résolution du handle impossible : {type(e).__name__}")
            return ""
        for v in vals:
            if v.get("type") == "URL":
                return (v.get("data") or {}).get("value", "")
        a.note(pass_no, source, url, f"handle sans URL (API HTTP {r.status_code})")
        return ""

    # -- sources -------------------------------------------------------------- #
    def from_elsevier(self, a, p):
        """API officielle Elsevier (ScienceDirect) : PDF des articles en accès libre."""
        pii = ""
        for u in a.landing_urls:
            pii = pii or extract_pii(u)
        doi_ok = a.doi and a.doi.lower().startswith(ELSEVIER_DOI_PREFIXES)
        if not (doi_ok or pii):
            return False
        url = (ELSEVIER_ARTICLE + "doi/" + quote(a.doi, safe="/()")) if doi_ok \
            else (ELSEVIER_ARTICLE + "pii/" + pii)
        ok, _, _ = self.try_url(a, p, "API Elsevier", url + "?httpAccept=application/pdf",
                                "elsevier", follow_html=False,
                                extra_headers={**self.http.elsevier_headers(),
                                               "Accept": "application/pdf"})
        last = a.attempts[-1][3] if a.attempts else ""
        if not ok and "AUTHENTICATION_ERROR" in last:
            # clé non habilitée à l'API de texte intégral : inutile d'insister pour les autres articles
            if not getattr(self, "_els_disabled", False):
                self._els_disabled = True
                self.http.els_denied = True
                log("⚠ API Elsevier : clé valide mais sans droits sur le texte intégral "
                    "(AUTHENTICATION_ERROR). Elsevier ne sert le texte intégral qu'aux requêtes venant du "
                    "réseau d'une institution abonnée, ou munies d'un jeton institutionnel "
                    "(--elsevier-insttoken). API désactivée pour le reste de l'exécution.")
            return False
        return ok

    def springer_meta_record(self, a, note):
        """Notice Springer Nature Meta API (v2) d'un DOI Springer ; mise en cache sur l'article.
        note(source, url, raison) consigne les échecs."""
        if a.springer_meta is not None:
            return a.springer_meta or None
        a.springer_meta = {}
        if getattr(self, "_springer_meta_disabled", False):
            return None
        src = "Springer Meta API"
        try:
            r = self.http.request("GET", SPRINGER_META, "springer", retries=2,
                                  params={"q": f'doi:"{a.doi}"', "api_key": self.http.springer_meta_key,
                                          "p": 1})
        except requests.RequestException as e:
            note(src, SPRINGER_META, f"erreur réseau : {type(e).__name__}")
            return None
        if r.status_code in (401, 403, 429):
            if not getattr(self, "_springer_meta_disabled", False):
                self._springer_meta_disabled = True
                why = "quota dépassé" if r.status_code == 429 else "clé refusée ou non habilitée"
                log(f"⚠ Springer Nature Meta API : {why} (HTTP {r.status_code}). Désactivée pour "
                    "le reste de l'exécution.")
            note(src, SPRINGER_META, f"API HTTP {r.status_code}")
            return None
        if r.status_code != 200:
            note(src, SPRINGER_META, f"API HTTP {r.status_code}")
            return None
        try:
            recs = r.json().get("records") or []
        except ValueError:
            note(src, SPRINGER_META, "réponse JSON invalide")
            return None
        rec = next((x for x in recs if (x.get("doi") or "").lower() == a.doi.lower()), None)
        if not rec:
            note(src, SPRINGER_META, "DOI inconnu de l'API Meta")
            return None
        a.springer_meta = rec
        return rec

    def from_springer_meta(self, a, p):
        """Lien PDF officiel donné par la Meta API pour les articles Springer Nature en libre accès."""
        src = "Springer Meta API"
        rec = self.springer_meta_record(a, lambda s, u, why: a.note(p, s, u, why))
        if not rec:
            return False
        if str(rec.get("openaccess", "")).lower() != "true":
            a.note(p, src, "", "article non libre d'après Springer Nature (openaccess=false)")
            return False
        a.oa_hint = True
        urls = [u.get("value") for u in (rec.get("url") or [])
                if isinstance(u, dict) and (u.get("format") or "").lower() == "pdf" and u.get("value")]
        if not urls:
            urls = [f"https://link.springer.com/content/pdf/{a.doi}.pdf"]
        for u in urls[:2]:
            ok, _, _ = self.try_url(a, p, src, u.replace("http://", "https://", 1), "publisher",
                                    follow_html=False)
            if ok:
                return True
        return False

    def from_s3(self, a, p):
        if not a.pmcid:
            return False
        src = "PMC S3"
        try:
            r = self.http.request("GET", S3_BASE, "s3",
                                  params={"list-type": "2", "prefix": f"{a.pmcid}.", "delimiter": "/"})
        except requests.RequestException as e:
            a.note(p, src, S3_BASE, f"erreur réseau : {type(e).__name__}")
            return False
        if r.status_code != 200:
            a.note(p, src, r.url, f"listing S3 HTTP {r.status_code}")
            return False
        vers = re.findall(r"<Prefix>(%s\.(\d+))/</Prefix>" % re.escape(a.pmcid), r.text)
        if not vers:
            a.note(p, src, r.url, "absent du jeu PMC Open Access (article PMC hors licence OA, "
                                  "sous embargo, ou manuscrit non diffusé)")
            return False
        pfx = max(vers, key=lambda v: int(v[1]))[0]
        pdf_url = f"{S3_BASE}{pfx}/{pfx}.pdf"
        try:
            m = self.http.request("GET", f"{S3_BASE}{pfx}/{pfx}.json", "s3")
            if m.status_code == 200:
                meta = m.json()
                u = meta.get("pdf_url") or ""
                if u.startswith("s3://pmc-oa-opendata/"):
                    pdf_url = S3_BASE + u[len("s3://pmc-oa-opendata/"):]
                elif u.startswith("http"):
                    pdf_url = u
                elif not u:
                    a.note(p, src, m.url, "métadonnées S3 : aucun PDF fourni par PMC pour cet article")
                    return False
        except (requests.RequestException, ValueError):
            pass
        ok, _, _ = self.try_url(a, p, src, pdf_url, "s3", follow_html=False)
        return ok

    def from_hal(self, a, p):
        """HAL (archive ouverte française) : PDF déposés, retrouvés par DOI via l'API officielle."""
        if not a.doi:
            return False
        src = "HAL"
        try:
            r = self.http.request("GET", HAL_API, "hal", retries=1,
                                  params={"q": f'doiId_s:"{a.doi}"', "wt": "json",
                                          "fl": "halId_s,fileMain_s,files_s,openAccess_bool"})
            docs = r.json().get("response", {}).get("docs", []) if r.status_code == 200 else []
        except (requests.RequestException, ValueError) as e:
            a.note(p, src, HAL_API, f"API HAL indisponible : {type(e).__name__}")
            return False
        files = []
        for d in docs:
            for u in [d.get("fileMain_s")] + list(d.get("files_s") or []):
                if u and u.lower().endswith(".pdf") and u not in files:
                    files.append(u)
        if not files:
            if docs:
                a.note(p, src, r.url, "notice HAL sans fichier PDF déposé")
            return False
        a.oa_hint = True
        for u in files[:2]:
            ok, _, _ = self.try_url(a, p, src, u, "repository", follow_html=False)
            if ok:
                return True
        return False

    def from_core(self, a, p):
        """CORE (core.ac.uk) : copies des dépôts institutionnels, souvent mises en cache par CORE."""
        if not (a.doi or len(a.title) > 20):
            return False
        # Budget de temps : le quota gratuit de CORE (~10 req/min, souvent moins) peut à lui
        # seul faire durer une recherche des dizaines de minutes pour un gain de 1 ou 2 PDF.
        if not hasattr(self, "_core_t0"):
            self._core_t0 = time.monotonic()
        budget = getattr(self.http, "core_budget", 300)
        if budget >= 0 and time.monotonic() - self._core_t0 > budget:
            if not getattr(self, "_core_disabled", False):
                self._core_disabled = True
                log(f"ℹ API CORE : budget de {budget // 60:.0f} min atteint, CORE désactivé pour "
                    "le reste de l'exécution (option --core-budget pour l'ajuster).")
            return False
        src = "CORE"
        q = f'doi:"{a.doi}"' if a.doi else f'title:"{a.title[:250]}"'
        try:
            r = self.http.request("GET", CORE_API, "core", retries=2,
                                  params={"q": q, "limit": 10},
                                  headers={"Authorization": f"Bearer {self.http.core_key}"})
        except requests.RequestException as e:
            a.note(p, src, CORE_API, f"erreur réseau : {type(e).__name__}")
            return False
        if r.status_code in (401, 403):
            if not getattr(self, "_core_disabled", False):
                self._core_disabled = True
                log(f"⚠ API CORE : clé refusée (HTTP {r.status_code}). CORE désactivé pour cette exécution.")
            a.note(p, src, CORE_API, f"clé API refusée (HTTP {r.status_code})")
            return False
        if r.status_code == 429:
            # quota dépassé : on ralentit les appels suivants (x2, plafond 60 s)
            with self.http._lock:
                old = self.http.intervals.get("core", 6.5)
                self.http.intervals["core"] = min(old * 2, 60.0)
            if self.http.intervals["core"] != old:
                log(f"⚠ API CORE : quota dépassé (429), intervalle porté à {self.http.intervals['core']:.0f} s.")
            if self.http.intervals["core"] >= 60.0 and not getattr(self, "_core_disabled", False):
                self._core_disabled = True
                log("⚠ API CORE : quota épuisé, CORE désactivé pour le reste de l'exécution.")
            a.note(p, src, CORE_API, "API HTTP 429 (quota CORE dépassé)")
            return False
        if r.status_code != 200:
            a.note(p, src, CORE_API, f"API HTTP {r.status_code}")
            return False
        try:
            results = r.json().get("results", []) or []
        except ValueError:
            a.note(p, src, CORE_API, "réponse JSON invalide")
            return False
        want_doi = (a.doi or "").lower()
        want_title = _norm(a.title)
        urls = []
        for w in results:
            wdoi = (w.get("doi") or "").lower()
            if want_doi:
                if wdoi and wdoi != want_doi:
                    continue
                if not wdoi and _norm(w.get("title") or "") != want_title:
                    continue
            elif _norm(w.get("title") or "") != want_title:
                continue
            cands = [w.get("downloadUrl")] + [l.get("url") for l in (w.get("links") or [])
                                              if l.get("type") == "download"]
            cands += list(w.get("sourceFulltextUrls") or [])
            for u in cands:
                if u and u not in urls:
                    urls.append(u)
        # copies hébergées par CORE d'abord (indépendantes de l'état du dépôt d'origine)
        urls.sort(key=lambda u: 0 if "core.ac.uk/download" in u else 1)
        if not urls:
            a.note(p, src, r.url, "aucune copie en texte intégral connue de CORE" if results
                   else "article inconnu de CORE")
            return False
        a.oa_hint = True
        for u in urls[:4]:
            svc = "core_dl" if "core.ac.uk" in u else "repository"
            ok, _, _ = self.try_url(a, p, src, u, svc, follow_html=(svc != "core_dl"))
            if ok:
                return True
        return False

    def from_unpaywall(self, a, p):
        if not a.doi:
            return [], []
        src = "Unpaywall"
        url = UNPAYWALL + quote(a.doi, safe="/()")
        try:
            r = self.http.request("GET", url, "unpaywall", params={"email": self.http.email})
        except requests.RequestException as e:
            a.note(p, src, url, f"erreur réseau : {type(e).__name__}")
            return [], []
        if r.status_code == 404:
            a.note(p, src, url, "DOI inconnu d'Unpaywall")
            return [], []
        if r.status_code != 200:
            a.note(p, src, url, f"API HTTP {r.status_code}")
            return [], []
        try:
            j = r.json()
        except ValueError:
            a.note(p, src, url, "réponse JSON invalide")
            return [], []
        if not j.get("is_oa"):
            a.note(p, src, url, "aucune version en libre accès connue d'Unpaywall")
            return [], []
        a.oa_hint = True
        locs = [j.get("best_oa_location")] + (j.get("oa_locations") or [])
        pdfs, lands, repos = [], [], []
        for loc in locs:
            if not loc:
                continue
            pdf_u = loc.get("url_for_pdf")
            land_u = loc.get("url_for_landing_page") or loc.get("url")
            if pdf_u and pdf_u not in pdfs:
                # copies des dépôts institutionnels en premier : rarement derrière Cloudflare
                (pdfs.insert(0, pdf_u) if loc.get("host_type") == "repository" else pdfs.append(pdf_u))
            if land_u and land_u != pdf_u:
                target = repos if loc.get("host_type") == "repository" else lands
                if land_u not in target:
                    target.append(land_u)
        a.repo_urls = set(repos) | {l.get("url_for_pdf") for l in locs
                                    if l and l.get("host_type") == "repository" and l.get("url_for_pdf")}
        lands = repos + lands
        if not pdfs:
            a.note(p, src, url, "version OA signalée mais sans URL PDF directe")
        return pdfs, lands

    def process(self, a, p):
        if STOP.is_set():
            return a
        a._tried = set()   # toutes les URL tentées pour cet article pendant cette passe
        tried = a._tried

        def attempt(src, u, service, follow=True):
            if not u:
                return False
            ok, html, final = self.try_url(a, p, src, u, service, follow_html=follow)
            return ok

        done = self.from_s3(a, p)
        if not done:
            for u in a.epmc_pdf_urls:
                svc = "europepmc" if "europepmc.org" in u else "publisher"
                if "europepmc.org" in u and "pdf=render" in u:
                    continue  # tenté plus bas
                if attempt("Europe PMC (lien PDF)", u, svc):
                    done = True
                    break
        lands = []
        if not done and a.is_springer and self.http.springer_meta_key \
                and not getattr(self, "_springer_meta_disabled", False):
            done = self.from_springer_meta(a, p)
        if not done and self.http.use("unpaywall"):
            pdfs, lands = self.from_unpaywall(a, p)
            if pdfs and not getattr(a, "manual_pdf", ""):
                a.manual_pdf = pdfs[0]   # lien PDF direct Unpaywall, pour le téléchargement manuel
            for u in pdfs:
                svc = "repository" if u in getattr(a, "repo_urls", ()) else "publisher"
                if attempt("Unpaywall", u, svc):
                    done = True
                    break
        # HAL et CORE : seulement pour les articles gratuits ou signalés en libre accès
        # (sur 776 articles, les interroger pour les non-gratuits coûtait ~1 h pour 0 PDF)
        if not done and a.free_like:
            done = self.from_hal(a, p)
        if not done and a.free_like and self.http.core_key and not getattr(self, "_core_disabled", False):
            done = self.from_core(a, p)
        if not done and self.http.elsevier_key and not getattr(self, "_els_disabled", False):
            done = self.from_elsevier(a, p)
        if not done:
            for u in a.openalex_pdf_urls:
                if attempt("OpenAlex", u, "publisher"):
                    done = True
                    break
        if not done:
            for u in [x for x in lands if x not in a.landing_urls][:3]:
                svc = "repository" if u in getattr(a, "repo_urls", ()) else "publisher"
                if attempt("Unpaywall (page du dépôt/éditeur)", u, svc):
                    done = True
                    break
        if not done and a.pmcid:
            done = attempt("Europe PMC (rendu PDF)",
                           f"https://europepmc.org/articles/{a.pmcid}?pdf=render", "europepmc",
                           follow=False)
        if not done:
            landing = []
            if a.doi:
                landing.append("https://doi.org/" + quote(a.doi, safe="/()"))
            landing += a.landing_urls
            for u in landing[:5]:
                if attempt("Site éditeur", u, "publisher"):
                    done = True
                    break
        if not done and a.doi.lower().startswith(KARGER_DOI_PREFIX) and a.free_like \
                and not any("Karger" in x[1] for x in a.attempts):
            a.note(p, "Karger", "", "Karger ne propose pas d'API publique de texte intégral "
                                    "(accès TDM sur contrat, par FTP) ; PDF à ouvrir à la main")
        if not done and not tried and not a.attempts:
            a.note(p, "—", "", "aucun identifiant exploitable (ni PMCID, ni DOI, ni lien éditeur)")
        elif not done and not any(x[0] == p for x in a.attempts):
            a.note(p, "—", "", "aucune URL de PDF trouvée par les sources interrogées")
        if done:
            self.source_ok[a.pdf_source.split(" (")[0]] += 1
        return a


# --------------------------------------------------------------------------- #
# Passe JSON : texte intégral structuré pour les articles sans PDF
# --------------------------------------------------------------------------- #
FULLTEXT_MIN_CHARS = 1500   # en deçà, le « texte intégral » n'est qu'un résumé


def _local(tag):
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _strip_ns(root):
    for el in root.iter():
        el.tag = _local(el.tag)
        for k in list(el.attrib):
            if "}" in k:
                el.attrib[_local(k)] = el.attrib.pop(k)
    return root


def _clean(s):
    return re.sub(r"\s+", " ", s or "").strip()


def new_doc(a, source, licence=""):
    return OrderedDict([
        ("schema", "pubmed_search-fulltext/1"), ("pmid", a.pmid), ("pmcid", a.pmcid), ("doi", a.doi),
        ("titre", a.title), ("journal", a.journal), ("annee", a.year), ("source", source),
        ("niveau", "résumé"), ("licence", licence),
        ("recupere_le", dt.datetime.now().strftime("%Y-%m-%dT%H:%M")),
        ("resume", ""), ("sections", []), ("figures", []), ("tableaux", []), ("references", []),
    ])


def _finish(doc):
    body = sum(len(s["texte"]) for s in doc["sections"])
    doc["niveau"] = "texte intégral" if body >= FULLTEXT_MIN_CHARS else "résumé"
    doc["nb_caracteres"] = body + len(doc["resume"])
    return doc


def jats_to_doc(article, a, source):
    """Article JATS (PMC, Europe PMC, Springer Nature) -> dictionnaire du schéma commun."""
    _strip_ns(article)
    lic = article.find(".//permissions/license")
    licence = ""
    if lic is not None:
        licence = lic.get("href", "") or _clean(" ".join(lic.itertext()))[:300]
    doc = new_doc(a, source, licence)
    t = article.find(".//article-meta/title-group/article-title")
    if t is not None and _clean("".join(t.itertext())):
        doc["titre"] = _clean("".join(t.itertext()))
    abs_el = article.find(".//article-meta/abstract")
    if abs_el is not None:
        doc["resume"] = "\n".join(_clean("".join(x.itertext())) for x in abs_el.iter("p")) \
            or _clean("".join(abs_el.itertext()))

    def text_of(el):
        # paragraphes de la section, sans les sous-sections ni les figures/tableaux
        out = []
        for ch in el:
            if ch.tag in ("p", "list", "disp-quote", "boxed-text", "statement"):
                out.append(_clean("".join(ch.itertext())))
        return "\n".join(x for x in out if x)

    def walk(sec, parents):
        title = _clean("".join(sec.find("title").itertext())) if sec.find("title") is not None else ""
        path = parents + ([title] if title else [])
        txt = text_of(sec)
        if txt:
            doc["sections"].append({"titre": " > ".join(path) or "(sans titre)",
                                    "type": sec.get("sec-type", ""), "texte": txt})
        for sub in sec.findall("sec"):
            walk(sub, path)

    body = article.find("body")
    if body is not None:
        loose = text_of(body)
        if loose:
            doc["sections"].append({"titre": "(texte)", "type": "", "texte": loose})
        for sec in body.findall("sec"):
            walk(sec, [])
    for f in article.iter("fig"):
        cap = f.find("caption")
        lab = _clean("".join(f.find("label").itertext())) if f.find("label") is not None else ""
        if cap is not None:
            doc["figures"].append({"label": lab, "legende": _clean("".join(cap.itertext()))})
    for tw in article.iter("table-wrap"):
        cap = tw.find("caption")
        lab = _clean("".join(tw.find("label").itertext())) if tw.find("label") is not None else ""
        doc["tableaux"].append({"label": lab, "legende": _clean("".join(cap.itertext())) if cap is not None else ""})
    for ref in article.iter("ref"):
        c = ref.find("mixed-citation")
        if c is None:
            c = ref.find("element-citation")
        if c is None:
            c = ref.find("citation")
        if c is not None:
            doc["references"].append(_clean(" ".join(c.itertext())))
    return _finish(doc)


def bioc_to_doc(data, a):
    """Réponse JSON de l'API BioC (NCBI) -> schéma commun."""
    colls = data if isinstance(data, list) else [data]
    passages = []
    for c in colls:
        for d in (c or {}).get("documents") or []:
            passages += d.get("passages") or []
    if not passages:
        return None
    licence = ""
    doc = new_doc(a, "NCBI BioC (PMC)")
    cur = None
    abstract = []
    for ps in passages:
        inf = ps.get("infons") or {}
        st = (inf.get("section_type") or "").upper()
        typ = (inf.get("type") or "").lower()
        txt = _clean(ps.get("text"))
        licence = licence or inf.get("license", "")
        if not txt:
            continue
        if st == "TITLE" or typ == "front":
            if typ in ("front", "title") and not doc["titre"]:
                doc["titre"] = txt
            continue
        if st == "ABSTRACT":
            if not typ.startswith("title"):
                abstract.append(txt)
            continue
        if st == "REF" or typ == "ref":
            if typ != "title":
                doc["references"].append(txt)
            continue
        if st == "FIG" or typ.startswith("fig"):
            doc["figures"].append({"label": inf.get("id", ""), "legende": txt})
            continue
        if st == "TABLE" or typ.startswith("table"):
            if "caption" in typ:
                doc["tableaux"].append({"label": inf.get("id", ""), "legende": txt})
            continue
        if typ.startswith("title"):
            cur = {"titre": txt, "type": st.lower(), "texte": ""}
            doc["sections"].append(cur)
            continue
        if cur is None or cur["type"] != st.lower():
            cur = {"titre": st.title() or "(texte)", "type": st.lower(), "texte": ""}
            doc["sections"].append(cur)
        cur["texte"] = (cur["texte"] + "\n" + txt).strip()
    doc["sections"] = [s for s in doc["sections"] if s["texte"]]
    doc["resume"] = "\n".join(abstract)
    doc["licence"] = licence
    return _finish(doc)


class JsonFetcher:
    """Récupère un JSON de texte intégral pour les articles restés sans PDF."""

    def __init__(self, http, json_dir, json_dirname, dl=None):
        self.http = http
        self.json_dir = json_dir
        self.json_dirname = json_dirname
        self.dl = dl   # Downloader : réutilise la notice Springer Meta mise en cache
        self.source_ok = Counter()
        self.level_ok = Counter()

    def _get(self, a, src, url, service, retries=2, **kw):
        try:
            r = self.http.request("GET", url, service, retries=retries, **kw)
        except requests.RequestException as e:
            a.jnote(src, url, f"erreur réseau : {type(e).__name__}")
            return None
        if r.status_code != 200:
            a.jnote(src, url, f"HTTP {r.status_code}")
            return None
        return r

    def save(self, a, doc):
        fname = a.pdf_filename()[:-4] + ".json"
        tmp = os.path.join(self.json_dir, fname + ".part")
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(doc, f, ensure_ascii=False, indent=1)
        os.replace(tmp, os.path.join(self.json_dir, fname))
        a.json_rel = f"{self.json_dirname}/{fname}"
        a.json_source = doc["source"]
        a.json_level = doc["niveau"]

    # -- sources ---------------------------------------------------------------- #
    def from_bioc(self, a):
        src = "NCBI BioC (PMC)"
        url = BIOC_PMC.format(a.pmcid or a.pmid)
        r = self._get(a, src, url, "bioc")
        if r is None:
            return None
        try:
            data = r.json()
        except ValueError:
            a.jnote(src, url, "absent du corpus BioC (article hors du sous-ensemble Open Access "
                              "de PMC et des manuscrits d'auteurs)")
            return None
        doc = bioc_to_doc(data, a)
        if doc is None:
            a.jnote(src, url, "réponse BioC vide")
        return doc

    def from_epmc(self, a):
        if not a.pmcid:
            return None
        src = "Europe PMC (JATS)"
        url = f"{EPMC_REST}{a.pmcid}/fullTextXML"
        try:
            # Europe PMC répond 404 ou 500 quand le texte intégral n'est pas dans son sous-ensemble
            # Open Access : pas de nouvel essai (sinon 2 attentes inutiles par article)
            r = self.http.request("GET", url, "europepmc", retries=0)
        except requests.RequestException as e:
            a.jnote(src, url, f"erreur réseau : {type(e).__name__}")
            return None
        if r.status_code in (404, 500):
            a.jnote(src, url, f"absent : texte intégral non diffusé par Europe PMC (HTTP {r.status_code})")
            return None
        if r.status_code != 200:
            a.jnote(src, url, f"HTTP {r.status_code}")
            return None
        try:
            root = ET.fromstring(r.content)
        except ET.ParseError:
            a.jnote(src, url, "XML illisible")
            return None
        art = root if _local(root.tag) == "article" else next(
            (e for e in root.iter() if _local(e.tag) == "article"), None)
        if art is None:
            a.jnote(src, url, "pas d'article JATS dans la réponse")
            return None
        return jats_to_doc(art, a, src)

    def from_springer_oa(self, a):
        if getattr(self, "_sn_oa_disabled", False):
            return None
        src = "Springer Nature Open Access API"
        try:
            r = self.http.request("GET", SPRINGER_OA_JATS, "springer", retries=2,
                                  params={"q": f'doi:"{a.doi}"', "api_key": self.http.springer_oa_key})
        except requests.RequestException as e:
            a.jnote(src, SPRINGER_OA_JATS, f"erreur réseau : {type(e).__name__}")
            return None
        if r.status_code in (401, 403, 429):
            if not getattr(self, "_sn_oa_disabled", False):
                self._sn_oa_disabled = True
                why = "quota dépassé" if r.status_code == 429 else "clé refusée ou non habilitée"
                log(f"⚠ Springer Nature Open Access API : {why} (HTTP {r.status_code}). "
                    "Désactivée pour le reste de l'exécution.")
            a.jnote(src, SPRINGER_OA_JATS, f"API HTTP {r.status_code}")
            return None
        if r.status_code != 200:
            a.jnote(src, SPRINGER_OA_JATS, f"API HTTP {r.status_code}")
            return None
        try:
            root = _strip_ns(ET.fromstring(r.content))
        except ET.ParseError:
            a.jnote(src, SPRINGER_OA_JATS, "XML illisible")
            return None
        art = next(root.iter("article"), None)
        if art is None:
            a.jnote(src, SPRINGER_OA_JATS, "absent de l'API Open Access (article non OA chez "
                                           "Springer Nature)")
            return None
        return jats_to_doc(art, a, src)

    def from_elsevier(self, a):
        if getattr(self, "_els_disabled", False) or getattr(self.http, "els_denied", False):
            return None
        src = "API Elsevier (JSON)"
        url = ELSEVIER_ARTICLE + "doi/" + quote(a.doi, safe="/()")
        r = self._get(a, src, url, "elsevier", params={"httpAccept": "application/json"},
                      headers={**self.http.elsevier_headers(), "Accept": "application/json"})
        if r is None:
            last = a.json_attempts[-1][2] if a.json_attempts else ""
            if last in ("HTTP 401", "HTTP 403"):
                self._els_disabled = True
                self.http.els_denied = True
            return None
        try:
            ftr = r.json().get("full-text-retrieval-response") or {}
        except ValueError:
            a.jnote(src, url, "réponse JSON invalide")
            return None
        core = ftr.get("coredata") or {}
        doc = new_doc(a, src, "openaccess" if str(core.get("openaccess")) in ("1", "true") else "")
        doc["titre"] = core.get("dc:title") or doc["titre"]
        doc["resume"] = _clean(core.get("dc:description") or "")
        txt = ftr.get("originalText")
        if isinstance(txt, dict):
            txt = json.dumps(txt, ensure_ascii=False)
        txt = (txt or "").strip()
        if txt:
            doc["sections"].append({"titre": "Texte intégral (format brut Elsevier)", "type": "",
                                    "texte": txt})
        doc = _finish(doc)
        if doc["niveau"] != "texte intégral":
            a.jnote(src, url, "résumé seulement (texte intégral non libre pour cette clé)")
        return doc

    def from_springer_meta(self, a):
        src = "Springer Meta API"
        rec = self.dl.springer_meta_record(a, a.jnote) if self.dl else None
        if not rec:
            return None
        doc = new_doc(a, src, "openaccess" if str(rec.get("openaccess")).lower() == "true" else "")
        doc["titre"] = rec.get("title") or doc["titre"]
        ab = rec.get("abstract")
        if isinstance(ab, dict):
            ps = ab.get("p")
            ab = "\n".join(ps) if isinstance(ps, list) else (ps or "")
        doc["resume"] = _clean(ab if isinstance(ab, str) else "")
        doc["mots_cles"] = rec.get("keyword") or rec.get("subjects") or []
        doc = _finish(doc)
        return doc if doc["resume"] else None

    def process(self, a):
        if STOP.is_set() or a.pdf_rel:
            return a
        candidates = []
        best_abstract = None
        steps = []
        if a.pmcid or a.free_like:
            steps.append(self.from_bioc)
        if a.pmcid:
            steps.append(self.from_epmc)
        if a.is_springer and self.http.springer_oa_key:
            steps.append(self.from_springer_oa)
        if a.is_elsevier and self.http.elsevier_key:
            steps.append(self.from_elsevier)
        if a.is_springer and self.http.springer_meta_key:
            steps.append(self.from_springer_meta)
        for step in steps:
            if STOP.is_set():
                break
            doc = step(a)
            if not doc:
                continue
            if doc["niveau"] == "texte intégral":
                candidates.append(doc)
                break
            if best_abstract is None and doc["resume"]:
                best_abstract = doc
        doc = candidates[0] if candidates else best_abstract
        if doc:
            self.save(a, doc)
            self.source_ok[doc["source"]] += 1
            self.level_ok[doc["niveau"]] += 1
        elif not steps:
            a.jnote("—", "", "aucune source JSON applicable (ni PMCID, ni gratuité, ni DOI "
                             "Springer/Elsevier avec clé)")
        if a.doi.lower().startswith(KARGER_DOI_PREFIX) and not (doc and doc["niveau"] == "texte intégral"):
            a.jnote("Karger", "", "pas d'API publique de texte intégral (TDM sur contrat, par FTP)")
        return a


# --------------------------------------------------------------------------- #
# Passe XML JATS : dossier « IA locale » (XML d'origine, Markdown, CSV, figures)
# --------------------------------------------------------------------------- #
JATS_IMG_EXT = (".jpg", ".jpeg", ".png", ".gif", ".tif", ".tiff", ".webp", ".bmp", ".svg")


def _s3_https(u):
    """s3://pmc-oa-opendata/KEY?md5=… -> https://pmc-oa-opendata.s3.amazonaws.com/KEY"""
    u = (u or "").split("?", 1)[0]
    if u.startswith("s3://pmc-oa-opendata/"):
        return S3_BASE + u[len("s3://pmc-oa-opendata/"):]
    return u


def _md_esc(s):
    return re.sub(r"\s+", " ", s or "").strip()


def _cell_md(s):
    return _md_esc(s).replace("|", "\\|")


class JatsMarkdown:
    """JATS (XML d'article PMC / Europe PMC) -> Markdown lisible par une IA, sans perte du contenu :
    sections, paragraphes, listes, tableaux complets (Markdown, ou HTML si cellules fusionnées),
    figures (image locale + légende), formules (LaTeX si fourni, sinon MathML), références."""

    def __init__(self, xml_bytes, media=None):
        self.root = _strip_ns(ET.fromstring(xml_bytes))
        self.media = media or {}         # nom (minuscules, sans extension) -> chemin relatif local
        self.online = {}                 # suppléments : nom (minuscules) -> URL en ligne (PMC S3)
        self.tables = []                 # [(label, nom de fichier CSV, grille)]
        self.figures = []                # [(label, chemin relatif ou "")]
        self.out = []

    # -- éléments en ligne ------------------------------------------------------------------- #
    def inl(self, el):
        parts = [el.text or ""]
        for c in el:
            t = c.tag
            inner = self.inl(c)
            if t in ("italic", "em"):
                s = f"*{inner.strip()}*" if inner.strip() else ""
            elif t == "bold":
                s = f"**{inner.strip()}**" if inner.strip() else ""
            elif t == "sup":
                s = f"<sup>{inner}</sup>"
            elif t == "sub":
                s = f"<sub>{inner}</sub>"
            elif t == "xref" and c.get("ref-type") == "bibr" and re.fullmatch(r"[\d,–\-\s]+", inner.strip() or "x") \
                    and not "".join(parts).rstrip().endswith(("[", "(")):
                s = f"[{inner.strip()}]"
            elif t == "ext-link" or t == "uri":
                href = c.get("href") or inner
                s = f"[{inner.strip() or href}]({href})" if href else inner
            elif t == "inline-formula":
                s = self.formula(c, inline=True)
            elif t in ("break",):
                s = " "
            elif t in ("fn", "table-wrap", "fig", "disp-formula", "list", "boxed-text"):
                s = ""          # éléments de bloc : rendus après le paragraphe
            else:
                s = inner
            parts.append(s)
            parts.append(c.tail or "")
        return "".join(parts)

    def formula(self, el, inline=False):
        tex = el.find(".//tex-math")
        if tex is not None and "".join(tex.itertext()).strip():
            body = "".join(tex.itertext()).strip()
            body = re.sub(r"^\\documentclass.*?\\begin\{document\}", "", body, flags=re.S)
            body = body.replace("\\end{document}", "").strip().strip("$")
            return f"${body}$" if inline else f"$$\n{body}\n$$"
        m = el.find(".//math")
        if m is not None:
            xml = ET.tostring(m, encoding="unicode")
            return f"`{_md_esc(''.join(m.itertext()))}`" if inline else f"```mathml\n{xml}\n```"
        return _md_esc("".join(el.itertext()))

    # -- blocs -------------------------------------------------------------------------------- #
    def blocks(self, el, level):
        for c in el:
            t = c.tag
            if t == "sec":
                title = c.find("title")
                label = c.findtext("label") or ""
                if title is not None or label:
                    head = _md_esc(f"{label} {self.inl(title) if title is not None else ''}")
                    self.out.append("#" * min(6, level) + " " + head)
                self.blocks(c, level + 1)
            elif t == "p":
                txt = _md_esc(self.inl(c))
                if txt:
                    self.out.append(txt)
                self.blocks_inside(c, level)
            elif t == "list":
                self.out.append("\n".join(self.list_(c)))
            elif t == "fig":
                self.fig(c)
            elif t == "fig-group":
                for f in c.findall("fig"):
                    self.fig(f)
            elif t == "table-wrap":
                self.table_wrap(c)
            elif t == "table-wrap-group":
                for tw in c.findall("table-wrap"):
                    self.table_wrap(tw)
            elif t == "disp-formula":
                self.out.append(self.formula(c))
            elif t == "boxed-text":
                self.out.append("> **Encadré** " + _md_esc(self.inl(c.find("caption"))) if c.find("caption") is not None
                                else "> **Encadré**")
                self.blocks(c, level + 1)
            elif t == "disp-quote":
                self.out.append("> " + _md_esc(" ".join(self.inl(p) for p in c.findall("p"))))
            elif t == "def-list":
                for d in c.findall("def-item"):
                    term = _md_esc(self.inl(d.find("term"))) if d.find("term") is not None else ""
                    dfn = _md_esc(" ".join(self.inl(p) for p in d.iter("p")))
                    self.out.append(f"- **{term}** : {dfn}")
            elif t == "supplementary-material":
                self.suppl(c)
            elif t in ("title", "label"):
                continue
            elif t in ("app", "ack", "glossary", "notes", "fn-group"):
                title = c.findtext("title") or {"ack": "Acknowledgments", "glossary": "Abbreviations",
                                                "fn-group": "Notes"}.get(t, "")
                if title:
                    self.out.append("#" * min(6, level) + " " + _md_esc(title))
                if t == "fn-group":
                    for fn in c.findall("fn"):
                        self.out.append("- " + _md_esc(" ".join(self.inl(p) for p in fn.findall("p"))))
                else:
                    self.blocks(c, level + 1)
            elif t in ("app-group",):
                self.blocks(c, level)
            elif len(c):
                self.blocks(c, level)

    def blocks_inside(self, p, level):
        for c in p:
            if c.tag in ("fig", "table-wrap", "disp-formula", "list", "boxed-text", "fig-group",
                         "table-wrap-group", "supplementary-material"):
                self.blocks(_wrap(c), level)

    def list_(self, el, depth=0):
        ordered = (el.get("list-type") or "") in ("order", "arabic", "alpha-lower", "alpha-upper",
                                                   "roman-lower", "roman-upper")
        lines = []
        for i, it in enumerate(el.findall("list-item"), 1):
            lab = it.findtext("label")
            mark = f"{lab}" if lab else (f"{i}." if ordered else "-")
            text = _md_esc(" ".join(self.inl(p) for p in it.findall("p")))
            lines.append("  " * depth + f"{mark} {text}")
            for sub in it.findall("list") + [x for p in it.findall("p") for x in p.findall("list")]:
                lines += self.list_(sub, depth + 1)
        return lines

    def _graphic(self, el):
        g = el.find(".//graphic")
        if g is None:
            return ""
        href = g.get("href") or ""
        key = os.path.splitext(os.path.basename(href))[0].lower()
        return self.media.get(href.lower()) or self.media.get(key) or ""

    def fig(self, f):
        label = _md_esc(f.findtext("label") or "Figure")
        cap = f.find("caption")
        captxt = _md_esc(" ".join(self.inl(x) for x in cap)) if cap is not None else ""
        path = self._graphic(f)
        self.figures.append((label, path))
        if path:
            self.out.append(f"![{label}]({path})")
        else:
            self.out.append(f"*[{label} : image non fournie avec le texte]*")
        self.out.append(f"**{label}.** {captxt}".strip())

    def suppl(self, s):
        label = _md_esc(s.findtext("label") or s.findtext("caption/title") or "Supplément")
        href = ""
        m = s.find(".//media")
        if m is not None:
            href = m.get("href") or ""
        local = self.media.get(href.lower()) or self.media.get(os.path.splitext(href)[0].lower()) or ""
        url = self.online.get(href.lower(), "")
        cap = _md_esc(" ".join(self.inl(p) for p in s.findall("caption/p")))
        link = (f" : [{os.path.basename(local)}]({local})" if local else
                f" : [{href}]({url}) (en ligne, non téléchargé)" if url else
                (f" : {href} (non téléchargé)" if href else ""))
        self.out.append(f"- **{label}**" + link + (f" — {cap}" if cap else ""))

    @staticmethod
    def grid(table):
        """<table> -> (grille [[texte]], lignes d'en-tête, fusion ?) ; les cellules fusionnées sont
        répétées dans la grille (lecture machine), la version HTML garde la fusion."""
        rows = []
        head_rows = 0
        for sect in ("thead", "tbody", "tfoot", None):
            trs = table.findall(f"{sect}/tr") if sect else table.findall("tr")
            if sect == "thead":
                head_rows = len(trs)
            rows += trs
        grid, pending, merged = [], {}, False
        for tr in rows:
            line, col = [], 0
            cells = [c for c in tr if c.tag in ("td", "th")]
            ci = 0
            while ci < len(cells) or col in pending:
                if col in pending:
                    txt, left = pending[col]
                    line.append(txt)
                    if left > 1:
                        pending[col] = (txt, left - 1)
                    else:
                        del pending[col]
                    col += 1
                    continue
                c = cells[ci]
                ci += 1
                txt = _md_esc("".join(c.itertext()))
                cs = int(re.sub(r"\D", "", c.get("colspan") or "1") or 1)
                rs = int(re.sub(r"\D", "", c.get("rowspan") or "1") or 1)
                if cs > 1 or rs > 1:
                    merged = True
                for k in range(cs):
                    line.append(txt)
                    if rs > 1:
                        pending[col] = (txt, rs - 1)
                    col += 1
            grid.append(line)
        width = max((len(r) for r in grid), default=0)
        grid = [r + [""] * (width - len(r)) for r in grid]
        return grid, head_rows, merged

    def table_html(self, table):
        out = ["<table>"]
        for tr in table.iter("tr"):
            cells = []
            for c in tr:
                if c.tag not in ("td", "th"):
                    continue
                attrs = "".join(f' {k}="{c.get(k)}"' for k in ("colspan", "rowspan") if c.get(k) not in (None, "1"))
                cells.append(f"<{c.tag}{attrs}>{htmlmod.escape(_md_esc(''.join(c.itertext())))}</{c.tag}>")
            out.append("<tr>" + "".join(cells) + "</tr>")
        out.append("</table>")
        return "\n".join(out)

    def table_wrap(self, tw):
        label = _md_esc(tw.findtext("label") or f"Table {len(self.tables) + 1}")
        cap = tw.find("caption")
        captxt = _md_esc(" ".join(self.inl(x) for x in cap)) if cap is not None else ""
        self.out.append(f"**{label}.** {captxt}".strip())
        table = tw.find(".//table")
        if table is None:
            path = self._graphic(tw)
            self.tables.append((label, "", None))
            self.out.append((f"![{label}]({path})\n" if path else "")
                            + "*Tableau fourni par l'éditeur sous forme d'image uniquement : contenu non "
                              "disponible en texte.*")
            return
        grid, head_rows, merged = self.grid(table)
        csv_name = re.sub(r"[^A-Za-z0-9]+", "_", label).strip("_") or f"Table_{len(self.tables) + 1}"
        names = {n for _, n, _ in self.tables}
        base, k = csv_name, 2
        while csv_name + ".csv" in names:
            csv_name = f"{base}_{k}"
            k += 1
        csv_name += ".csv"
        self.tables.append((label, csv_name, grid))
        if grid and not merged and head_rows <= 1:
            w = len(grid[0])
            lines = ["| " + " | ".join(_cell_md(x) for x in grid[0]) + " |", "|" + "---|" * w]
            lines += ["| " + " | ".join(_cell_md(x) for x in r) + " |" for r in grid[1:]]
            self.out.append("\n".join(lines))
        elif grid:
            self.out.append(self.table_html(table))
        self.out.append(f"*Données : tableaux/{csv_name}*")
        foot = tw.find("table-wrap-foot")
        if foot is not None:
            notes = [_md_esc(self.inl(x)) for x in foot.iter() if x.tag in ("p",)]
            if notes:
                self.out.append("*Notes : " + " ".join(notes) + "*")

    # -- article ----------------------------------------------------------------------------- #
    def render(self, meta_line=""):
        r = self.root
        am = r.find(".//article-meta")
        title = _md_esc(self.inl(am.find("title-group/article-title"))) if am is not None and \
            am.find("title-group/article-title") is not None else "(sans titre)"
        self.out.append(f"# {title}")
        if am is not None:
            names = []
            for c in am.findall("contrib-group/contrib"):
                if (c.get("contrib-type") or "author") != "author":
                    continue
                sn, gn = c.findtext("name/surname") or "", c.findtext("name/given-names") or ""
                coll = _md_esc("".join(c.find("collab").itertext())) if c.find("collab") is not None else ""
                n = _md_esc(f"{gn} {sn}") or coll
                if n:
                    names.append(n)
            if names:
                self.out.append("**Auteurs :** " + ", ".join(names))
        if meta_line:
            self.out.append(meta_line)
        lic = r.find(".//permissions/license")
        if lic is not None:
            href = lic.get("href") or ""
            self.out.append("**Licence :** " + (_md_esc("".join(lic.itertext()))[:300] or href)
                            + (f" ({href})" if href and href not in "".join(lic.itertext()) else ""))
        if am is not None:
            kws = [_md_esc("".join(k.itertext())) for k in am.findall("kwd-group/kwd")]
            if kws:
                self.out.append("**Mots-clés :** " + "; ".join(k for k in kws if k))
            for ab in am.findall("abstract"):
                kind = ab.get("abstract-type") or ""
                self.out.append("## " + ("Abstract" if not kind else f"Abstract ({kind})"))
                if ab.find("sec") is not None:
                    self.blocks(ab, 3)
                else:
                    for p in ab.findall("p"):
                        self.out.append(_md_esc(self.inl(p)))
        body = r.find("body")
        if body is not None:
            self.blocks(body, 2)
        back = r.find("back")
        if back is not None:
            self.blocks(back, 2)
            refs = back.findall(".//ref-list/ref")
            if refs:
                self.out.append("## References")
                for i, ref in enumerate(refs, 1):
                    lab = ref.findtext("label") or str(i)
                    cit = next((x for x in (ref.find("mixed-citation"), ref.find("element-citation"),
                                            ref.find("citation")) if x is not None), ref)
                    if cit.tag == "mixed-citation":
                        txt = _md_esc(re.sub(r"\s+([.,;:])", r"\1", "".join(cit.itertext())))
                    else:
                        txt = _md_esc(", ".join(x.strip() for x in cit.itertext() if x.strip()))
                    ids = [f"{p.get('pub-id-type')}: {p.text}" for p in cit.findall("pub-id") if p.text]
                    self.out.append(f"{lab.strip('.').strip()}. {txt}" + (f" [{'; '.join(ids)}]" if ids and
                                                                         not all(i.split(': ')[1] in txt for i in ids) else ""))
        fg = r.find("floats-group")
        if fg is not None:
            self.blocks(fg, 2)
        md = "\n\n".join(x for x in self.out if x and x.strip())
        return re.sub(r"\n{3,}", "\n\n", md) + "\n"


def _wrap(el):
    w = ET.Element("wrap")
    w.append(el)
    return w


def write_table_csv(path, grid):
    import csv
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(grid)


class JatsFetcher:
    """XML JATS des articles PMC (Open Access et manuscrits d'auteurs) : PMC S3 (XML + figures +
    suppléments), sinon Europe PMC (XML seul) ; puis Markdown + CSV des tableaux."""

    def __init__(self, http, jats_dir, jats_dirname):
        self.http = http
        self.dir = jats_dir
        self.dirname = jats_dirname
        self.source_ok = Counter()
        self.n_fig = 0
        self.n_tab = 0
        self.n_tab_img = 0

    def _get(self, url, service, **kw):
        try:
            r = self.http.request("GET", url, service, **kw)
        except requests.RequestException as e:
            return None, f"erreur réseau : {type(e).__name__}"
        if r.status_code != 200:
            return None, f"HTTP {r.status_code}"
        return r, ""

    def s3_meta(self, a):
        r, why = self._get(S3_BASE, "s3", params={"list-type": "2", "prefix": f"{a.pmcid}.", "delimiter": "/"})
        if r is None:
            return None, f"PMC S3 : listing {why}"
        vers = re.findall(r"<Prefix>(%s\.(\d+))/</Prefix>" % re.escape(a.pmcid), r.text)
        if not vers:
            return None, "PMC S3 : absent du jeu PMC (hors Open Access / manuscrits d'auteurs)"
        pfx = max(vers, key=lambda v: int(v[1]))[0]
        m, why = self._get(f"{S3_BASE}{pfx}/{pfx}.json", "s3")
        meta = {}
        if m is not None:
            try:
                meta = m.json()
            except ValueError:
                meta = {}
        meta.setdefault("xml_url", f"{S3_BASE}{pfx}/{pfx}.xml")
        meta["_prefix"] = pfx
        return meta, ""

    def _download(self, url, dest, max_bytes=None):
        try:
            r = self.http.request("GET", url, "s3", stream=True)
        except requests.RequestException:
            return False
        if r.status_code != 200:
            r.close()
            return False
        size = int(r.headers.get("Content-Length") or 0)
        if max_bytes and size > max_bytes:
            r.close()
            return False
        tmp = dest + ".part"
        n = 0
        with open(tmp, "wb") as f:
            for chunk in r.iter_content(65536):
                n += len(chunk)
                if max_bytes and n > max_bytes:
                    f.close()
                    os.remove(tmp)
                    r.close()
                    return False
                f.write(chunk)
        os.replace(tmp, dest)
        return True

    def process(self, a):
        if not a.pmcid or STOP.is_set():
            if not a.pmcid:
                a.jats_note = "pas de PMCID : pas de XML JATS librement diffusé"
            return a
        base = a.pdf_filename()[:-4]
        folder = os.path.join(self.dir, base)
        xml, source, media, licence, notes = None, "", {}, "", []
        meta, why = self.s3_meta(a)
        if meta:
            r, why2 = self._get(_s3_https(meta.get("xml_url")), "s3")
            if r is not None:
                xml, source = r.content, "PMC S3"
                licence = meta.get("license_code") or ""
            else:
                notes.append(f"PMC S3 : XML {why2}")
        else:
            notes.append(why)
        if xml is None:
            url = f"{EPMC_REST}{a.pmcid}/fullTextXML"
            r, why3 = self._get(url, "europepmc", retries=0)
            if r is not None and r.content.strip().startswith(b"<"):
                xml, source = r.content, "Europe PMC"
            else:
                notes.append(f"Europe PMC : {why3 or 'réponse vide'}")
        if xml is None:
            a.jats_note = " ; ".join(notes)
            return a
        try:
            conv = JatsMarkdown(xml)
        except ET.ParseError as e:
            a.jats_note = f"{source} : XML illisible ({e})"
            return a
        os.makedirs(folder, exist_ok=True)
        with open(os.path.join(folder, base + ".xml"), "wb") as f:
            f.write(xml)
        n_img = n_sup = 0
        online = {}
        if meta and source == "PMC S3":
            for u in meta.get("media_urls") or []:
                if STOP.is_set():
                    break
                url = _s3_https(u)
                name = unquote(os.path.basename(urlparse(url).path))
                if not name:
                    continue
                if not name.lower().endswith(JATS_IMG_EXT):
                    # suppléments (tableurs, vidéos…) : lien en ligne seulement (volume souvent important)
                    online[name.lower()] = url
                    n_sup += 1
                    continue
                os.makedirs(os.path.join(folder, "figures"), exist_ok=True)
                rel = f"figures/{name}"
                if self._download(url, os.path.join(folder, "figures", name)):
                    media[name.lower()] = rel
                    media.setdefault(os.path.splitext(name)[0].lower(), rel)
                    n_img += 1
                else:
                    notes.append(f"image non téléchargée : {name}")
        elif source == "Europe PMC":
            notes.append("Europe PMC fournit le XML sans les images : figures non téléchargées")
        conv.media = media
        conv.online = online
        meta_line = " · ".join(x for x in (
            a.journal, a.year, f"PMID {a.pmid}", a.pmcid, f"DOI {a.doi}" if a.doi else "",
            f"source : {source}" + (f" (licence {licence})" if licence else "")) if x)
        md = conv.render(f"*{meta_line}*")
        with open(os.path.join(folder, base + ".md"), "w", encoding="utf-8") as f:
            f.write(md)
        nt = 0
        for label, csv_name, grid in conv.tables:
            if grid is None:
                self.n_tab_img += 1
                continue
            os.makedirs(os.path.join(folder, "tableaux"), exist_ok=True)
            write_table_csv(os.path.join(folder, "tableaux", csv_name), grid)
            nt += 1
        with self.http._lock:
            self.source_ok[source] += 1
            self.n_fig += n_img
            self.n_tab += nt
        a.jats_rel = f"{self.dirname}/{base}/{base}.md"
        a.jats_source = source
        a.jats_info = (f"XML + Markdown, {n_img} figure(s), {nt} tableau(x) CSV"
                       + (f", {n_sup} supplément(s) en lien" if n_sup else "")
                       + (f", {sum(1 for t in conv.tables if t[2] is None)} tableau(x) en image"
                          if any(t[2] is None for t in conv.tables) else ""))
        a.jats_note = " ; ".join(n for n in notes if not n.startswith(("PMC S3 : absent", "PMC S3 : listing")))
        return a


JATS_README = """Dossier « IA locale » — PubMed Search
====================================

Un sous-dossier par article (PMID-Auteur(Année)), avec :
  • <article>.xml  : XML JATS d'origine (PMC / Europe PMC), complet et sans perte ;
  • <article>.md   : le même article en Markdown (titres, texte, tableaux complets, légendes,
                     formules, références) : le format à donner à une IA (modèle local, RAG…) ;
  • tableaux/      : un fichier CSV par tableau (cellules fusionnées répétées), pour R, Excel, Python ;
  • figures/       : les images des figures (référencées dans le Markdown).
Les fichiers supplémentaires ne sont pas téléchargés (volume souvent important) : le Markdown
donne leur lien PMC.

Limites : seuls les articles présents dans PMC (sous-ensemble Open Access et manuscrits d'auteurs)
ont un XML librement diffusé. Europe PMC fournit parfois le XML sans les figures. Les tableaux
déposés en image par l'éditeur restent des images. Les graphiques sont des images : les valeurs
exactes se trouvent dans les tableaux ou les suppléments quand les auteurs les ont publiées.
"""


# --------------------------------------------------------------------------- #
# Mots-clés et termes MeSH équivalents (une ligne = des synonymes)
# --------------------------------------------------------------------------- #
MESH_LOOKUP = "https://id.nlm.nih.gov/mesh/lookup/descriptor"
MESH_DETAILS = "https://id.nlm.nih.gov/mesh/lookup/details"


def kw_norm(s):
    s = unicodedata.normalize("NFKC", s or "").lower()
    s = re.sub(r"[‐‑‒–—\-/_]+", " ", s)
    s = re.sub(r"[\"'“”‘’«»*]", "", s)
    s = re.sub(r"\s+", " ", s).strip(" .,;:")
    return s


def _valid_sf(sf):
    if not 2 <= len(sf) <= 10 or len(sf.split()) > 2:
        return False
    if not re.search(r"[A-Za-z]", sf) or not sf[0].isalnum():
        return False
    if re.search(r"[=%;,<>±]", sf) or sf.lower() in ("ie", "eg", "i.e", "e.g", "vs", "ref", "refs"):
        return False
    if re.fullmatch(r"[a-z]+", sf) and len(sf) > 3:   # mot ordinaire entre parenthèses
        return False
    if re.fullmatch(r"\d+(\.\d+)?\s*[a-z]{0,3}", sf, re.I):   # « 12 h », « 3.5 mg »
        return False
    return True


def _best_long_form(sf, cand):
    """Algorithme de Schwartz & Hearst (2003) : la forme longue se termine juste avant la
    parenthèse ; chaque caractère de l'abréviation y est retrouvé de droite à gauche, et le 1er
    au début d'un mot."""
    s, l = len(sf) - 1, len(cand) - 1
    sfl, lfl = sf.lower(), cand.lower()
    while s >= 0:
        c = sfl[s]
        if not c.isalnum():
            s -= 1
            continue
        while (l >= 0 and lfl[l] != c) or (s == 0 and l > 0 and lfl[l - 1].isalnum()):
            l -= 1
        if l < 0:
            return None
        l -= 1
        s -= 1
    start = cand.rfind(" ", 0, l + 1) + 1
    lf = cand[start:].strip(" ,;:-")
    if len(lf) <= len(sf) or sf.lower() in lf.lower().split():
        return None
    return lf


def find_abbreviations(text):
    out = []
    for m in re.finditer(r"\(([^()\[\]]{2,10})\)", text or ""):
        sf = m.group(1).strip()
        if not _valid_sf(sf):
            continue
        before = re.split(r"(?<=[.;:!?])\s|[\(\)\[\]]", text[:m.start()])[-1]
        words = before.split()
        maxw = min(len(sf) + 5, len(sf) * 2)
        lf = _best_long_form(sf, " ".join(words[-maxw:]))
        if lf and len(lf.split()) <= maxw:
            out.append((sf, lf))
    return out


class _UF:
    def __init__(self):
        self.p = {}

    def find(self, x):
        self.p.setdefault(x, x)
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, a, b):
        ra, rb = self.find(a), self.find(b)
        if ra != rb:
            self.p[rb] = ra


class KeywordGroups:
    """Regroupe termes MeSH, mots-clés d'auteurs et abréviations en lignes de synonymes."""

    def __init__(self, http, quiet=False):
        self.http = http
        self.uf = _UF()
        self.display = {}           # clé normalisée -> Counter des graphies
        self.arts = {}              # clé -> set(PMID) (ou d'indices de requête)
        self.is_mesh = set()        # clés qui sont des descripteurs MeSH
        self.mesh_ui = {}           # nom MeSH (clé) -> UI (D…)
        self.entry_done = set()
        self.ambiguous = {}         # abréviation -> [formes longues]
        self.quiet = quiet
        self.n_entry = 0
        self.n_atm = 0
        self.n_abbr = 0
        self.n_fail = 0

    def add(self, term, src_id, mesh=False, ui=""):
        k = kw_norm(term)
        if not k or len(k) < 2:
            return ""
        self.uf.find(k)
        self.display.setdefault(k, Counter())[term.strip()] += 1
        if src_id is not None:
            self.arts.setdefault(k, set()).add(src_id)
        else:
            self.arts.setdefault(k, set())
        if mesh:
            self.is_mesh.add(k)
            if ui:
                self.mesh_ui[k] = ui
        return k

    # -- synonymes officiels du MeSH (entry terms) -------------------------------------------- #
    def entry_terms(self, uis):
        """esummary db=mesh (lots de 200) : DS_MeshTerms = descripteur + synonymes (entry terms)."""
        uis = [u for u in dict.fromkeys(uis) if u and u not in self.entry_done and re.fullmatch(r"D\d+", u)]
        for batch in chunks(uis, 200):
            if STOP.is_set():
                return
            ids = ["68" + u[1:] for u in batch]
            try:
                r = self.http.request("POST", EUTILS + "esummary.fcgi", "ncbi",
                                      data=self.http.ncbi(db="mesh", id=",".join(ids), retmode="json"))
                r.raise_for_status()
                res = r.json().get("result", {})
            except Exception as e:  # noqa
                self.n_fail += 1
                log(f"  ⚠ MeSH (synonymes) : {short_err(e)}")
                continue
            rev = {}
            for k, v in self.mesh_ui.items():
                rev.setdefault(v, k)
            for u, i in zip(batch, ids):
                self.entry_done.add(u)
                doc = res.get(i) or {}
                terms = doc.get("ds_meshterms") or []
                if not isinstance(terms, list) or not terms:
                    continue
                head = kw_norm(terms[0])
                key = rev.get(u, head)
                for t in terms:
                    if ", " in t:          # formes permutées (« Traps, Extracellular »)
                        continue
                    k = self.add(t, None)
                    if k:
                        self.uf.union(key, k)
                        self.n_entry += 1

    def mesh_ui_by_name(self, name):
        try:
            r = self.http.request("GET", MESH_LOOKUP, "meshapi", retries=1,
                                  params={"label": name, "match": "exact", "limit": 1})
            j = r.json() if r.status_code == 200 else []
        except Exception:  # noqa
            self.n_fail += 1
            return ""
        if j and isinstance(j, list):
            return str(j[0].get("resource", "")).rsplit("/", 1)[-1]
        return ""

    # -- correspondance automatique de PubMed (Automatic Term Mapping) ------------------------ #
    def atm(self, term):
        """MeSH auxquels PubMed rattache un terme libre (translationset d'ESearch)."""
        try:
            r = self.http.request("POST", EUTILS + "esearch.fcgi", "ncbi", retries=1,
                                  data=self.http.ncbi(db="pubmed", term=term, retmax=0, retmode="json"))
            r.raise_for_status()
            ts = r.json().get("esearchresult", {}).get("translationset", []) or []
        except Exception:  # noqa
            self.n_fail += 1
            return []
        out = []
        for t in ts:
            for m in re.findall(r'"([^"]+)"\[MeSH Terms\]', t.get("to", "")):
                if m not in out:
                    out.append(m)
        return out

    def map_free_terms(self, keys, max_lookups):
        done = 0
        for k in keys:
            if STOP.is_set() or done >= max_lookups:
                break
            disp = self.best(k)
            # abréviations courtes : le mapping PubMed se trompe souvent (NLR -> NLR Proteins) ;
            # elles sont rattachées par la détection dans les résumés
            if re.fullmatch(r"[A-Z0-9\-]{2,6}s?", disp):
                continue
            done += 1
            meshes = self.atm(disp)
            if len(meshes) != 1:
                continue
            mk = self.add(meshes[0], None, mesh=True)
            if mk and mk != k:
                self.uf.union(mk, k)
                self.n_atm += 1
                if mk not in self.mesh_ui:
                    ui = self.mesh_ui_by_name(meshes[0]) or self.mesh_ui_by_name(meshes[0].title())
                    if ui:
                        self.mesh_ui[mk] = ui
            if done % 50 == 0 and not self.quiet:
                log(f"  mots-clés : {done} rattachés au MeSH…")

    def best(self, k):
        c = self.display.get(k)
        if not c:
            return k
        # graphie la plus fréquente ; à égalité, celle qui a des majuscules (sigles, noms propres)
        return sorted(c.items(), key=lambda kv: (-kv[1], kv[0].islower(), kv[0]))[0][0]

    def add_abbreviations(self, texts, keep_all_min=2):
        """texts : [(id, texte)] ; paires (sigle, forme longue) gardées si vues ≥ keep_all_min fois
        ou si l'un des deux termes est déjà un mot-clé ou un MeSH."""
        pairs = Counter()
        disp_lf = {}
        where = {}
        for sid, txt in texts:
            for sf, lf in find_abbreviations(txt):
                key = (sf, kw_norm(lf))
                pairs[key] += 1
                disp_lf.setdefault(key, Counter())[lf] += 1
                where.setdefault(key, set()).add(sid)
        by_sf = {}
        for (sf, lfk), n in pairs.items():
            if n >= keep_all_min or lfk in self.arts or kw_norm(sf) in self.arts:
                by_sf.setdefault(sf, []).append((n, lfk))
        def jacc(x, y):
            a, b = set(x.split()), set(y.split())
            return len(a & b) / max(1, len(a | b))
        for sf, lst in by_sf.items():
            lst.sort(reverse=True)
            main = lst[0][1]
            # variantes d'écriture d'un même sens (« neutrophil-to-lymphocyte » / « neutrophil
            # lymphocyte ») : rattachées ; sens différents (NLR = NOD-like receptor) : seul le plus
            # fréquent est rattaché au sigle, les autres sont signalés
            variants = [lfk for _, lfk in lst if lfk == main or jacc(lfk, main) >= 0.5]
            others = [lfk for _, lfk in lst if lfk not in variants]
            for lfk in variants:
                lf = disp_lf[(sf, lfk)].most_common(1)[0][0]
                for sid in where[(sf, lfk)]:
                    k_lf = self.add(lf, sid)
                    k_sf = self.add(sf, sid)
                if k_lf and k_sf:
                    self.uf.union(k_sf, k_lf)
                    self.n_abbr += 1
            if others:
                self.ambiguous[sf] = [disp_lf[(sf, g)].most_common(1)[0][0] for g in [main] + others[:3]]

    def groups(self):
        g = {}
        for k in self.display:
            g.setdefault(self.uf.find(k), []).append(k)
        out = []
        for keys in g.values():
            ids = set()
            for k in keys:
                ids |= self.arts.get(k, set())
            if not ids:
                continue
            keys.sort(key=lambda k: (-len(self.arts.get(k, ())), k not in self.is_mesh, k))
            seen, terms = set(), []
            for k in keys:
                d = self.best(k)
                if kw_norm(d) not in seen:
                    seen.add(kw_norm(d))
                    terms.append(d)
            out.append((len(ids), terms, any(k in self.is_mesh for k in keys), keys))
        out.sort(key=lambda x: (-x[0], x[1][0].lower()))
        return out

    def write(self, path, title_lines, n_items, unit="articles", sources=None):
        groups = self.groups()
        de_unit = ("d'" if unit[:1] in "aeiouéè" else "de ") + unit
        with open(path, "w", encoding="utf-8") as f:
            for t in title_lines:
                f.write(f"# {t}\n")
            f.write("# Une ligne = des termes équivalents séparés par « = » ; lignes triées par nombre "
                    f"{de_unit} (sur {n_items}), le terme le plus employé en premier.\n")
            f.write("# Sources : " + (sources or "descripteurs MeSH et leurs synonymes officiels (entry terms), "
                    "mots-clés des auteurs, correspondance automatique de PubMed, sigles définis dans les "
                    "titres et résumés (méthode de Schwartz & Hearst, 2003).") + "\n#\n")
            for n, terms, _, _ in groups:
                f.write(" = ".join(terms) + "\n")
            if self.ambiguous:
                f.write("#\n# Sigles à sens multiples (rattachés au sens le plus fréquent) :\n")
                for sf, lfs in sorted(self.ambiguous.items()):
                    f.write(f"#   {sf} : " + " | ".join(lfs) + "\n")
        return groups


def build_keywords(http, articles, path, query, max_lookups=150, offline=False):
    kg = KeywordGroups(http)
    for a in articles:
        for ui, name in a.mesh:
            kg.add(name, a.pmid, mesh=True, ui=ui)
        for kw in a.keywords:
            kg.add(kw, a.pmid)
    n_terms = len(kg.display)
    if not offline:
        log(f"  • synonymes MeSH ({len(set(kg.mesh_ui.values()))} descripteurs)…")
        kg.entry_terms(list(kg.mesh_ui.values()))
        free = sorted((k for k in kg.display if k not in kg.is_mesh
                       and kg.uf.find(k) == k and len(kg.arts.get(k, ())) >= 1),
                      key=lambda k: -len(kg.arts.get(k, ())))
        log(f"  • rattachement des mots-clés d'auteurs au MeSH (≤ {max_lookups} termes)…")
        kg.map_free_terms(free, max_lookups)
        kg.entry_terms(list(kg.mesh_ui.values()))
    log("  • sigles définis dans les titres et résumés…")
    kg.add_abbreviations([(a.pmid, f"{a.title} {a.abstract}") for a in articles])
    groups = kg.write(path, [f"Mots-clés et termes MeSH — requête : {query}",
                             f"{len(articles)} articles, {n_terms} termes distincts au départ"],
                      len(articles))
    log(f"  {len(groups)} ligne(s) de termes équivalents ({kg.n_entry} synonymes MeSH, {kg.n_atm} "
        f"rattachements PubMed, {kg.n_abbr} sigles)")
    if kg.n_fail:
        log(f"  ⚠ {kg.n_fail} requête(s) MeSH/PubMed sans réponse : synonymes incomplets.")
    return groups


QUERY_SKIP_TAGS = ("pt", "dp", "la", "sb", "filter", "lang", "publication type", "date", "edat", "mhda",
                   "crdt", "pdat", "sex", "ptyp", "jour", "ta", "au", "1au", "lastau", "pmid", "uid",
                   "doi", "aid", "vi", "ip", "pg", "volume", "issue", "subset", "language", "mesh subheading", "sh")


def query_terms(q):
    """Termes d'une requête PubMed (sans opérateurs, filtres de date/type/langue ni étiquettes)."""
    out = []
    q = re.sub(r"\b(AND|OR|NOT)\b", "\x00", q)
    for part in re.split(r"[\x00()]", q):
        part = part.strip()
        if not part:
            continue
        m = re.match(r'^(.*?)\s*\[([^\]]+)\]\s*$', part)
        tag = ""
        if m:
            part, tag = m.group(1).strip(), m.group(2).strip().lower()
            tag = re.sub(r":.*$", "", tag)
        if tag in QUERY_SKIP_TAGS or re.fullmatch(r'"?\s*\d{4}([/\-:]\d{1,4})*\s*"?.*', part):
            continue
        term = part.strip('" ').rstrip("*").strip()
        if len(term) < 2 or re.fullmatch(r"last \d+ (years?|days?|months?)", term, re.I):
            continue
        out.append((term, tag))
    return out


def run_query_keywords(args):
    """--keywords-from-queries FICHIER : termes des requêtes d'une liste (une requête par ligne)."""
    with open(args.keywords_from_queries, encoding="utf-8") as f:
        queries = [ln.strip() for ln in f if ln.strip() and not ln.startswith("#")]
    http = Http(args.email, args.api_key, sources=[])
    kg = KeywordGroups(http, quiet=True)
    log(f"Termes des requêtes : {len(queries)} requête(s)")
    pending_ui = []
    for qi, q in enumerate(queries, 1):
        for term, tag in query_terms(q):
            mesh = tag in ("mh", "mesh", "mesh terms", "majr", "mesh major topic", "mh:noexp")
            k = kg.add(term, qi, mesh=mesh)
            if not k:
                continue
            if mesh:
                ui = kg.mesh_ui_by_name(term) or kg.mesh_ui_by_name(term.title())
                if ui:
                    kg.mesh_ui[k] = ui
                    pending_ui.append(ui)
            elif not re.fullmatch(r"[A-Z0-9\-]{2,6}s?", term):
                for m in kg.atm(term)[:3]:
                    mk = kg.add(m, qi, mesh=True)
                    kg.uf.union(k, mk)
                    kg.n_atm += 1
                    ui = kg.mesh_ui_by_name(m.title()) or kg.mesh_ui_by_name(m)
                    if ui:
                        kg.mesh_ui[mk] = ui
                        pending_ui.append(ui)
        log(f"  [{qi}/{len(queries)}] {q[:90]}")
    kg.entry_terms(pending_ui)
    out = args.output or os.path.splitext(args.keywords_from_queries)[0] + "_mots-cles.txt"
    label = os.path.basename(args.keywords_from_queries).replace(" (requêtes).txt", "")
    groups = kg.write(out, [f"Termes des requêtes de la liste : {label}",
                            f"{len(queries)} requête(s)"], len(queries), unit="requêtes",
                      sources="termes des requêtes, correspondance automatique de PubMed vers le MeSH, "
                              "synonymes officiels du MeSH (entry terms).")
    if kg.n_fail:
        log(f"⚠ {kg.n_fail} requête(s) MeSH/PubMed sans réponse : synonymes incomplets (réseau ?).")
    log(f"{len(groups)} ligne(s) de termes équivalents → {out}")
    log(f"KEYWORDS_FILE: {out}")
    return 0


# --------------------------------------------------------------------------- #
# Sorties
# --------------------------------------------------------------------------- #
def write_excel(path, articles, summary_rows, kw_groups=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Articles"
    headers = ["PMID", "DOI", "Titre de l'article", "Type d'article", "1er auteur",
               "Dernier auteur", "Autres auteurs", "Email (auteurs)", "Gratuit",
               "PDF (lien local)", "PMCID", "Année", "Journal", "Source du PDF", "Tri / rang",
               "JSON (lien local)", "Source du JSON", "Contenu du JSON",
               "IA locale : Markdown (lien local)", "Contenu JATS"]
    ws.append(headers)
    hfill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hfill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for a in articles:
        row = [a.pmid, a.doi, a.title, "; ".join(a.pub_types), a.first_author, a.last_author,
               a.other_authors, a.emails, a.free, a.pdf_rel or "non téléchargé", a.pmcid,
               a.year, a.journal, a.pdf_source, a.sort_label,
               a.json_rel or ("" if a.pdf_rel else "-"), a.json_source, a.json_level,
               a.jats_rel or (a.jats_note[:120] if a.jats_note else ""), a.jats_info]
        ws.append([clean_xl(v) for v in row])
        r = ws.max_row
        ws.cell(r, 1).hyperlink = f"https://pubmed.ncbi.nlm.nih.gov/{a.pmid}/"
        ws.cell(r, 1).font = Font(color="0563C1", underline="single")
        if a.doi:
            ws.cell(r, 2).hyperlink = "https://doi.org/" + a.doi
        if a.pdf_rel:
            ws.cell(r, 10).hyperlink = a.pdf_rel
            ws.cell(r, 10).font = Font(color="0563C1", underline="single")
        else:
            ws.cell(r, 10).font = Font(color="C00000")
        if a.json_rel:
            ws.cell(r, 16).hyperlink = a.json_rel
            ws.cell(r, 16).font = Font(color="0563C1", underline="single")
        if a.jats_rel:
            ws.cell(r, 19).hyperlink = a.jats_rel
            ws.cell(r, 19).font = Font(color="0563C1", underline="single")
    widths = [11, 26, 60, 24, 18, 18, 40, 40, 17, 42, 13, 7, 30, 22, 26, 42, 26, 16, 46, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Résumé")
    ws2.append(["Paramètre", "Valeur"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hfill
    for k, v in summary_rows:
        ws2.append([clean_xl(k), clean_xl(v)])
        if v == "" and k:
            ws2.cell(ws2.max_row, 1).font = Font(bold=True)
    ws2.column_dimensions["A"].width = 48
    ws2.column_dimensions["B"].width = 100
    for row in ws2.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    if kw_groups:
        ws3 = wb.create_sheet("Mots-clés")
        ws3.append(["Termes équivalents", "Articles", "MeSH"])
        for c in ws3[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = hfill
        for n, terms, is_mesh, _ in kw_groups:
            ws3.append([clean_xl(" = ".join(terms)), n, "oui" if is_mesh else ""])
        ws3.column_dimensions["A"].width = 110
        ws3.column_dimensions["B"].width = 10
        ws3.column_dimensions["C"].width = 8
        ws3.freeze_panes = "A2"
        ws3.auto_filter.ref = ws3.dimensions
    wb.save(path)


def classify(reason):
    r = reason.lower()
    if "anti-robot" in r or "preuve de travail" in r:
        return "Protection anti-robot (Cloudflare/CAPTCHA/JS)"
    if r.startswith("http 403"):
        return "HTTP 403 (accès refusé)"
    if r.startswith("http 404") or "listing s3 http 404" in r:
        return "HTTP 404 (introuvable)"
    if r.startswith("http "):
        return "Autre erreur HTTP"
    if r.startswith("ignoré"):
        return "Ignoré (site web PMC : téléchargement automatisé interdit/protégé)"
    if "absent du jeu pmc" in r:
        return "Absent du PMC Open Access (S3)"
    if "html reçu" in r:
        return "Page HTML au lieu d'un PDF"
    if "timeout" in r or "réseau" in r or "interrompu" in r:
        return "Réseau / délai dépassé"
    if "aucune version" in r or "inconnu d'unpaywall" in r:
        return "Pas de version OA connue (Unpaywall)"
    if "aucun" in r:
        return "Aucune URL / identifiant"
    return "Autre"


def advice(articles, http, getfree):
    """Conseils concrets d'après les échecs (clés manquantes, articles hors Open Access…)."""
    failed = [a for a in articles if not a.pdf_rel and a.free_like]
    out = []
    for svc in sorted({k for k, _ in getattr(http, "paused_log", [])}):
        n = http.skipped.get(svc, 0)
        out.append(f"{BREAKER_SERVICES.get(svc, svc)} ne répondait pas pendant la recherche (erreurs serveur "
                   f"répétées) : service mis en pause, {n} requête(s) évitée(s). Il s'agit d'une panne "
                   "côté serveur, pas d'une erreur de la requête : relancer plus tard la même recherche "
                   "permettra d'obtenir les liens manquants.")
    if not getfree or not failed:
        return out
    if not any(a.attempts for a in articles):
        out.append("Aucun téléchargement n'a été tenté (recherche interrompue avant la passe 1) : "
                   "relancer la recherche.")
        return out
    failed = [a for a in failed if a.attempts]
    els = [a for a in failed if a.is_elsevier]
    if els and not http.elsevier_key:
        out.append(f"{len(els)} article(s) libre(s) Elsevier/Lancet/Cell bloqué(s) par Cloudflare : une clé "
                   "API Elsevier gratuite (dev.elsevier.com) permet souvent de les obtenir par l'API officielle "
                   "(PDF puis JSON). PMID : " + ", ".join(a.pmid for a in els))
    elif els and getattr(http, "els_denied", False):
        out.append(f"{len(els)} article(s) Elsevier non obtenu(s) : la clé est valide mais Elsevier refuse le "
                   "texte intégral hors abonnement (AUTHENTICATION_ERROR). Solutions : lancer la recherche "
                   "depuis le réseau d'une institution abonnée à ScienceDirect, ou demander à sa bibliothèque "
                   "un jeton institutionnel (--elsevier-insttoken). Sinon : onglet PDF → Importer les PDF d'une recherche. PMID : "
                   + ", ".join(a.pmid for a in els))
    elif els:
        out.append(f"{len(els)} article(s) Elsevier non obtenu(s) malgré la clé : voir le détail. PMID : "
                   + ", ".join(a.pmid for a in els))
    sn = [a for a in failed if a.is_springer]
    if sn and not http.springer_meta_key:
        out.append(f"{len(sn)} article(s) Springer Nature non obtenu(s) : une clé Meta API (et Open Access "
                   "API pour le JSON) peut aider. PMID : " + ", ".join(a.pmid for a in sn))
    pmc = [a for a in failed if a.pmcid and any("absent du jeu PMC" in x[3] for x in a.attempts)]
    if pmc:
        out.append(f"{len(pmc)} article(s) lisible(s) gratuitement sur PMC mais hors du sous-ensemble Open "
                   "Access (aucune voie automatisée autorisée) : utiliser l'onglet PDF → Importer les PDF d'une recherche. PMID : "
                   + ", ".join(a.pmid for a in pmc))
    return out


def _write_jats_section(f, articles):
    pmc = [a for a in articles if a.pmcid]
    miss = [a for a in pmc if not a.jats_rel]
    part = [a for a in pmc if a.jats_rel and a.jats_note]
    f.write("\n" + "-" * 78 + "\nPASSE XML JATS (IA locale)\n" + "-" * 78 + "\n")
    f.write(f"Articles avec PMCID : {len(pmc)} | XML obtenus : {len(pmc) - len(miss)} | "
            f"sans XML : {len(miss)} | sans PMCID (pas de XML libre) : "
            f"{sum(1 for a in articles if not a.pmcid)}\n")
    for a in miss:
        f.write(f"  PMID {a.pmid} | {a.pmcid} : {a.jats_note or 'non traité (arrêt)'}\n")
    for a in part:
        f.write(f"  PMID {a.pmid} | {a.pmcid} (obtenu) : {a.jats_note}\n")


def write_error_file(path, articles, meta_lines, getfree, tips=(), getjats=False):
    failed = [a for a in articles if not a.pdf_rel and a.free_like]
    other = [a for a in articles if not a.pdf_rel and not a.free_like]
    cats = Counter()
    for a in failed:
        last_pass = max((x[0] for x in a.attempts), default=0)
        for x in a.attempts:
            if x[0] == last_pass:
                cats[classify(x[3])] += 1
    with open(path, "w", encoding="utf-8") as f:
        f.write("pubmed_search.py — rapport des échecs de téléchargement\n")
        f.write("=" * 78 + "\n")
        for line in meta_lines:
            f.write(line + "\n")
        f.write("\n")
        if not getfree:
            f.write("Option --getfreepaper non activée : aucun téléchargement de PDF tenté.\n")
            if getjats:
                _write_jats_section(f, articles)
            return
        f.write(f"Articles gratuits/OA non téléchargés : {len(failed)}\n")
        f.write(f"Articles sans mention de gratuité ni version OA détectée, non téléchargés : "
                f"{len(other)}\n\n")
        if tips:
            f.write("CONSEILS :\n")
            for t in tips:
                f.write(f"  • {t}\n")
            f.write("\n")
        f.write("Synthèse des causes (dernière passe, une ligne par tentative) :\n")
        for k, v in cats.most_common():
            f.write(f"  {v:5d}  {k}\n")
        f.write("\n" + "-" * 78 + "\nDÉTAIL — ARTICLES GRATUITS / OA NON TÉLÉCHARGÉS\n" + "-" * 78 + "\n")
        for a in failed:
            f.write(f"\nPMID {a.pmid} | DOI {a.doi or '-'} | PMCID {a.pmcid or '-'} | "
                    f"Gratuit : {a.free}{' + OA détecté' if a.oa_hint else ''}\n")
            f.write(f"  Titre : {a.title[:150]}\n")
            manual = getattr(a, "manual_pdf", "") or (f"https://pmc.ncbi.nlm.nih.gov/articles/{a.pmcid}/" if a.pmcid
                      else (f"https://doi.org/{a.doi}" if a.doi else f"https://pubmed.ncbi.nlm.nih.gov/{a.pmid}/"))
            f.write(f"  À ouvrir manuellement : {manual}\n")
            for p, src, url, reason in a.attempts:
                f.write(f"  [passe {p}] {src} : {reason}" + (f"\n             URL : {url}" if url else "") + "\n")
        if other:
            f.write("\n" + "-" * 78 + "\nARTICLES NON GRATUITS (aucune version libre détectée) — "
                    "PMID : " + ", ".join(a.pmid for a in other) + "\n")
        jtried = [a for a in articles if not a.pdf_rel and (a.json_attempts or a.json_rel)]
        if jtried:
            jc = Counter()
            for a in jtried:
                for src, _, why in a.json_attempts:
                    jc[f"{src} : {why.split(' (')[0] if 'absent' in why else classify(why)}"] += 1
            f.write("\n" + "-" * 78 + "\nPASSE JSON (articles sans PDF)\n" + "-" * 78 + "\n")
            f.write(f"JSON texte intégral : {sum(1 for a in jtried if a.json_level == 'texte intégral')} | "
                    f"JSON résumé seulement : {sum(1 for a in jtried if a.json_level == 'résumé')} | "
                    f"sans JSON : {sum(1 for a in jtried if not a.json_rel)}\n")
            f.write("Synthèse des échecs JSON :\n")
            for k, v in jc.most_common():
                f.write(f"  {v:5d}  {k}\n")
            for a in jtried:
                res = f"{a.json_level} ({a.json_source})" if a.json_rel else "aucun JSON"
                f.write(f"\nPMID {a.pmid} | DOI {a.doi or '-'} | PMCID {a.pmcid or '-'} → {res}\n")
                for src, url, why in a.json_attempts:
                    f.write(f"  [JSON] {src} : {why}" + (f"\n             URL : {url}" if url else "") + "\n")
        if getjats:
            _write_jats_section(f, articles)


# --------------------------------------------------------------------------- #
# Mode identification : dossier de PDF quelconque -> PDF renommés, Excel, Word
# (« recherche inversée » : on part du PDF pour retrouver la notice PubMed)
# --------------------------------------------------------------------------- #
CROSSREF = "https://api.crossref.org/works"
DOI_RE = re.compile(r"\b(10\.\d{4,9}/[^\s\"<>{}|\\^`]+)", re.I)
PMID_TXT_RE = re.compile(r"\bPMID\s*:?\s*(\d{6,9})\b", re.I)
MONTHS = {str(i): m for i, m in enumerate(
    ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])}
BAD_TITLES = ("untitled", "microsoft word", "document", "article", "pdf", "manuscript", "title")


def _clean_doi(d):
    d = unquote(d or "").strip().rstrip(".,;:)]}>'\"")
    d = re.sub(r"(\.pdf|\.full|/full|/abstract|/epdf|/pdf)$", "", d, flags=re.I)
    if d.count("(") < d.count(")"):          # parenthèse fermante collée au DOI
        d = d[:d.rfind(")")]
    return d


def _words(s):
    return [w for w in re.findall(r"[a-z0-9]+", unicodedata.normalize("NFKD", s or "")
                                  .encode("ascii", "ignore").decode().lower()) if len(w) > 3]


def _containment(title, text):
    """Part des mots significatifs du titre présents dans le texte du PDF (0 à 1)."""
    tw = _words(title)
    if not tw:
        return 0.0
    txt = set(_words(text))
    return sum(1 for w in tw if w in txt) / len(tw)


def _similar(a, b):
    na, nb = " ".join(_words(a)), " ".join(_words(b))
    if not (na and nb):
        return 0.0
    try:
        import difflib
        return difflib.SequenceMatcher(None, na, nb).ratio()
    except ImportError:   # Python intégré d'une ancienne version de l'application
        wa, wb = set(na.split()), set(nb.split())
        return 2 * len(wa & wb) / (len(wa) + len(wb))


class PdfInfo:
    """Ce qu'on peut lire d'un PDF sans réseau : texte, métadonnées, titre probable."""

    def __init__(self, path):
        self.path = path
        self.name = os.path.basename(path)
        self.text1 = ""        # page 1
        self.text = ""         # pages 1-2
        self.meta_doi = ""
        self.meta_title = ""
        self.big_title = ""    # texte en plus gros caractères de la page 1
        self.error = ""
        if PdfReader is None:
            self.error = "module pypdf absent"
            return
        try:
            r = PdfReader(path)
            if r.is_encrypted:
                try:
                    r.decrypt("")
                except Exception:  # noqa
                    self.error = "PDF protégé par mot de passe"
                    return
            pages = r.pages[:2]
            frags = []

            def visit(t, cm, tm, font, size):
                if t and t.strip():
                    sc = max(abs(tm[0]), abs(tm[3]), 1e-6) * max(abs(cm[0]), abs(cm[3]), 1e-6)
                    frags.append((round((size or 0) * sc, 1), t))
            if pages:
                try:
                    self.text1 = pages[0].extract_text(visitor_text=visit) or ""
                except Exception:  # noqa
                    self.text1 = pages[0].extract_text() or ""
            self.text = self.text1 + " " + " ".join((pg.extract_text() or "") for pg in pages[1:])
            # titre = fragments de la plus grande taille (hors lettrines d'un caractère)
            sizes = sorted({sz for sz, t in frags if len(t.strip()) > 3}, reverse=True)
            for sz in sizes[:3]:
                cand = re.sub(r"\s+", " ", " ".join(t for s2, t in frags if abs(s2 - sz) < 0.6)).strip()
                if 20 <= len(cand) <= 350 and len(cand.split()) >= 4:
                    self.big_title = cand
                    break
            md = r.metadata or {}
            raw = " ".join(str(v) for v in md.values() if v)
            try:
                if r.xmp_metadata is not None:
                    raw += " " + r.xmp_metadata.stream.get_data().decode("utf-8", "ignore")
            except Exception:  # noqa
                pass
            m = DOI_RE.search(raw)
            if m:
                self.meta_doi = _clean_doi(m.group(1))
            t = str(md.get("/Title") or "").strip()
            if len(t) >= 20 and not t.lower().startswith(BAD_TITLES) and not t.lower().endswith(
                    (".doc", ".docx", ".pdf", ".tex", ".indd")):
                self.meta_title = t
        except Exception as e:  # noqa
            self.error = f"PDF illisible ({type(e).__name__})"

    def filename_doi(self):
        fname = unquote(self.name)
        if re.match(r"10\.\d{4,9}_", fname):        # « 10.1000_xyz.pdf » : / remplacé par _
            fname = fname.replace("_", "/", 1)
        m = DOI_RE.search(fname)
        return _clean_doi(re.sub(r"\.pdf$", "", m.group(1), flags=re.I)) if m else ""

    def doi_candidates(self):
        """DOI dans l'ordre de confiance : nom du fichier, métadonnées, page 1 (près du
        mot « doi »), puis pages 1-2."""
        out = []

        def add(d):
            d = _clean_doi(d)
            if d and d.lower() not in [x.lower() for x in out]:
                out.append(d)
        if self.filename_doi():
            add(self.filename_doi())
        if self.meta_doi:
            add(self.meta_doi)
        flat = re.sub(r"\s+", " ", self.text1)
        for m in re.finditer(r"(?:doi(?:\.org)?\s*[:/]?\s*|dx\.doi\.org/)(10\.\d{4,9}/\S+)", flat, re.I):
            add(m.group(1))
        for m in DOI_RE.finditer(flat):
            add(m.group(1))
        for m in DOI_RE.finditer(re.sub(r"\s+", " ", self.text)):
            add(m.group(1))
        return out[:8]

    def pmid_candidates(self):
        out = []
        m = re.match(r"^(?:PMID[\s_-]*)?(\d{6,9})(?=[-_\s(.])", self.name, re.I)
        if m:
            out.append(m.group(1))
        out += [x for x in PMID_TXT_RE.findall(self.text1) if x not in out]
        return out

    def title_candidates(self):
        return [t for t in (self.big_title, self.meta_title) if t]


def _crossref_article(item):
    """Notice Crossref -> Article (sans PMID)."""
    a = Article("")
    a.doi = item.get("DOI", "")
    a.title = re.sub(r"<[^>]+>", "", " ".join(item.get("title") or [])).strip()
    a.journal = " ".join(item.get("container-title") or [])
    a.journal_abbr = " ".join(item.get("short-container-title") or []) or a.journal
    parts = ((item.get("issued") or item.get("published-print") or item.get("published-online")
              or {}).get("date-parts") or [[None]])[0]
    a.year = str(parts[0]) if parts and parts[0] else ""
    a.month = MONTHS.get(str(parts[1]), "") if len(parts) > 1 and parts[1] else ""
    a.volume = item.get("volume", "")
    a.issue = item.get("issue", "")
    a.pages = item.get("page", "") or item.get("article-number", "")
    a.pub_types = [item.get("type", "")] if item.get("type") else []
    for au in item.get("author") or []:
        last = au.get("family") or au.get("name") or ""
        given = au.get("given") or ""
        ini = "".join(p[0] for p in re.split(r"[\s\-.]+", given) if p and p[0].isalpha()).upper()
        if last:
            a.authors.append({"name": f"{last} {ini}".strip(), "last": last, "initials": ini,
                              "emails": []})
    return a


class Identifier:
    def __init__(self, http):
        self.http = http
        self.cache = {}

    # -- PubMed -------------------------------------------------------------- #
    def pubmed_record(self, pmid):
        if pmid not in self.cache:
            try:
                self.cache[pmid] = efetch_articles(self.http, [pmid], quiet=True).get(pmid)
            except Exception:  # noqa
                self.cache[pmid] = None
        return self.cache[pmid]

    def pmid_for_doi(self, doi):
        try:
            _, ids, _ = esearch(self.http, f'"{doi}"[aid]', "relevance", 3)
        except Exception:  # noqa
            return None
        return ids[0] if len(ids) == 1 else None

    def pubmed_by_title(self, title):
        words = [w for w in re.findall(r"[A-Za-z0-9]+", title) if len(w) > 3][:12]
        if len(words) < 3:
            return None, 0.0
        try:
            _, ids, _ = esearch(self.http, " AND ".join(f"{w}[ti]" for w in words), "relevance", 5)
        except Exception:  # noqa
            return None, 0.0
        if not ids:
            return None, 0.0
        try:
            recs = efetch_articles(self.http, ids, quiet=True)
        except Exception:  # noqa
            return None, 0.0
        best, score = None, 0.0
        for rec in recs.values():
            sc = _similar(title, rec.title)
            if sc > score:
                best, score = rec, sc
        if best is not None:
            self.cache[best.pmid] = best
        return best, score

    # -- Crossref ------------------------------------------------------------ #
    def crossref_doi(self, doi):
        try:
            r = self.http.request("GET", f"{CROSSREF}/{quote(doi, safe='/')}", "crossref", retries=1,
                                  params={"mailto": self.http.email})
            if r.status_code != 200:
                return None
            return _crossref_article(r.json().get("message") or {})
        except Exception:  # noqa
            return None

    def crossref_title(self, title):
        try:
            r = self.http.request("GET", CROSSREF, "crossref", retries=1,
                                  params={"query.bibliographic": title[:300], "rows": 3,
                                          "mailto": self.http.email})
            if r.status_code != 200:
                return None, 0.0
            items = (r.json().get("message") or {}).get("items") or []
        except Exception:  # noqa
            return None, 0.0
        best, score = None, 0.0
        for it in items:
            a = _crossref_article(it)
            sc = _similar(title, a.title)
            if sc > score:
                best, score = a, sc
        return best, score

    # -- un PDF ---------------------------------------------------------------- #
    def identify(self, info):
        """Renvoie (article ou None, méthode, fiabilité 'sûre'/'à vérifier', remarque)."""
        has_text = len(_words(info.text)) >= 30
        fname_doi = info.filename_doi()
        # 1. PMID explicite (nom du fichier ou « PMID: » imprimé)
        for pmid in info.pmid_candidates():
            rec = self.pubmed_record(pmid)
            if rec and (not has_text or _containment(rec.title, info.text) >= 0.6):
                return rec, "PMID (nom du fichier ou texte)", "sûre", ""
        # 2. DOI (nom, métadonnées, texte) -> PubMed, sinon Crossref
        doubtful = None
        for doi in info.doi_candidates():
            pmid = self.pmid_for_doi(doi)
            rec = self.pubmed_record(pmid) if pmid else None
            src = "DOI → PubMed"
            if rec is None:
                rec = self.crossref_doi(doi)
                src = "DOI → Crossref (absent de PubMed)"
            if rec is None:
                continue
            cont = _containment(rec.title, info.text) if has_text else None
            if cont is None and fname_doi and doi.lower() == fname_doi.lower():
                return rec, src + " (DOI du nom de fichier)", "sûre", ""
            if cont is None or cont >= 0.6:
                note = "" if cont is not None else "texte du PDF illisible : titre non vérifié"
                return rec, src, "sûre" if cont is not None else "à vérifier", note
            if cont >= 0.3 and doubtful is None:
                doubtful = (rec, src, "à vérifier",
                            f"titre retrouvé à {cont:.0%} dans le PDF (DOI peut-être celui "
                            "d'une référence citée)")
        # 3. titre -> PubMed, puis Crossref
        for title in info.title_candidates():
            rec, sc = self.pubmed_by_title(title)
            if rec is not None and sc >= 0.90:
                return rec, "Titre → PubMed", "sûre", ""
            if rec is not None and sc >= 0.75 and doubtful is None:
                doubtful = (rec, "Titre → PubMed", "à vérifier", f"titre proche à {sc:.0%}")
            rec, sc = self.crossref_title(title)
            if rec is not None and sc >= 0.90:
                pmid = self.pmid_for_doi(rec.doi) if rec.doi else None
                prec = self.pubmed_record(pmid) if pmid else None
                if prec:
                    return prec, "Titre → Crossref → PubMed", "sûre", ""
                return rec, "Titre → Crossref (absent de PubMed)", "sûre", ""
            if rec is not None and sc >= 0.75 and doubtful is None:
                doubtful = (rec, "Titre → Crossref", "à vérifier", f"titre proche à {sc:.0%}")
        if doubtful:
            return doubtful
        if info.error:
            return None, "", "", info.error
        if not has_text:
            return None, "", "", ("pas de texte lisible (PDF scanné ?) ni de DOI dans le nom ou les "
                                  "propriétés du fichier")
        return None, "", "", "ni DOI, ni PMID, ni titre retrouvé dans PubMed ou Crossref"


def ident_filename(a):
    """Nom du PDF : PMID-Auteur_Initiales(Année).pdf, ou NoPMID-… hors PubMed."""
    name = a.pdf_filename()
    if not a.pmid:
        name = "NoPMID" + name
    return name


def vancouver(a):
    """Référence au style Vancouver (NLM, Citing Medicine) : 6 auteurs puis et al."""
    au = [x["name"] for x in a.authors]
    if len(au) > 6:
        authors = ", ".join(au[:6]) + ", et al."
    elif au:
        authors = ", ".join(au) + "."
    else:
        authors = ""
    title = (a.title or "").strip()
    if title and title[-1] not in ".?!":
        title += "."
    jour = (a.journal_abbr or a.journal or "").strip()
    mon = MONTHS.get(a.month.lstrip("0"), a.month) if a.month else ""
    date = " ".join(x for x in (a.year, mon) if x)
    loc = ""
    if a.volume:
        loc += a.volume
    if a.issue:
        loc += f"({a.issue})"
    if a.pages:
        loc += f":{a.pages}"
    src = f"{jour}. " if jour else ""
    src += date + (f";{loc}" if loc else "") + "."
    ids = []
    if a.doi:
        ids.append(f"doi:{a.doi}.")
    pm = []
    if a.pmid:
        pm.append(f"PMID: {a.pmid}")
    if a.pmcid:
        pm.append(f"PMCID: {a.pmcid}")
    if pm:
        ids.append("; ".join(pm) + ".")
    return " ".join(x for x in (authors, title, src, " ".join(ids)) if x)


def _xml_esc(s):
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def write_docx(path, title, subtitle, refs):
    """Document Word minimal (WordprocessingML) sans dépendance : titre, sous-titre et
    références numérotées avec retrait suspendu."""
    import zipfile
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

    def para(text, style=None, num=None):
        ppr = f'<w:pStyle w:val="{style}"/>' if style else ""
        if num is not None:
            ppr += ('<w:tabs><w:tab w:val="left" w:pos="567"/></w:tabs>'
                    '<w:spacing w:after="120"/><w:ind w:left="567" w:hanging="567"/>')
            runs = f'<w:r><w:t>{num}.</w:t></w:r><w:r><w:tab/></w:r>'
        else:
            runs = ""
        runs += f'<w:r><w:t xml:space="preserve">{_xml_esc(text)}</w:t></w:r>'
        return f"<w:p><w:pPr>{ppr}</w:pPr>{runs}</w:p>"
    body = para(title, "Title") + para(subtitle, "Subtitle")
    body += "".join(para(r, num=i) for i, r in enumerate(refs, 1))
    doc = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>{body}'
           '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1418" w:right="1418" '
           'w:bottom="1418" w:left="1418" w:header="709" w:footer="709" w:gutter="0"/></w:sectPr>'
           '</w:body></w:document>')
    styles = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles {W}>'
              '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" '
              'w:cs="Calibri" w:eastAsia="Calibri"/><w:sz w:val="22"/><w:lang w:val="fr-FR"/></w:rPr>'
              '</w:rPrDefault><w:pPrDefault><w:pPr><w:spacing w:after="80" w:line="264" '
              'w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>'
              '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
              '<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/>'
              '<w:pPr><w:spacing w:after="60"/></w:pPr><w:rPr><w:b/><w:sz w:val="36"/></w:rPr></w:style>'
              '<w:style w:type="paragraph" w:styleId="Subtitle"><w:name w:val="Subtitle"/>'
              '<w:basedOn w:val="Normal"/><w:pPr><w:spacing w:after="240"/></w:pPr>'
              '<w:rPr><w:color w:val="595959"/><w:sz w:val="20"/></w:rPr></w:style></w:styles>')
    ctypes = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types '
              'xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
              '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
              '<Default Extension="xml" ContentType="application/xml"/>'
              '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-'
              'officedocument.wordprocessingml.document.main+xml"/>'
              '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-'
              'officedocument.wordprocessingml.styles+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships '
            'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/officeDocument" Target="word/document.xml"/></Relationships>')
    drels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships '
             'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
             '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
             'relationships/styles" Target="styles.xml"/></Relationships>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ctypes)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", doc)
        z.writestr("word/styles.xml", styles)
        z.writestr("word/_rels/document.xml.rels", drels)


def _sort_key(a):
    last = a.authors[0]["last"] if a.authors else (a.title or "")
    last = unicodedata.normalize("NFKD", last).encode("ascii", "ignore").decode().lower()
    return (last, a.year or "9999", (a.title or "").lower())


def run_identify(src_dir, output, outdir, http):
    t0 = time.monotonic()
    src_dir = os.path.abspath(os.path.expanduser(src_dir))
    if not os.path.isdir(src_dir):
        sys.exit(f"Dossier introuvable : {src_dir}")
    stem = ascii_slug(os.path.splitext(os.path.basename(output or "Bibliographie"))[0],
                      keep=r"A-Za-z0-9_\-") or "Bibliographie"
    now = dt.datetime.now()
    base = os.path.abspath(os.path.expanduser(outdir)) if outdir else src_dir
    folder = os.path.join(base, f"{stem}_{now:%d.%m.%Y_%Hh%M}")
    k = 2
    while os.path.exists(folder):
        folder = os.path.join(base, f"{stem}_{now:%d.%m.%Y_%Hh%M}_{k}")
        k += 1
    pdf_dirname = f"PDF_{stem}"
    pdf_dir = os.path.join(folder, pdf_dirname)
    check_dir = os.path.join(pdf_dir, "A_verifier")
    pdfs = sorted(p for p in glob.glob(os.path.join(src_dir, "*"))
                  if p.lower().endswith(".pdf") and os.path.isfile(p))
    log(f"pubmed_search.py v{VERSION} — identification de PDF — {now:%d/%m/%Y %Hh%M}")
    log(f"Dossier analysé : {src_dir}  ({len(pdfs)} PDF)")
    log(f"Résultats       : {folder}")
    if not pdfs:
        sys.exit("Aucun fichier PDF dans ce dossier.")
    if PdfReader is None:
        log("⚠ Module 'pypdf' absent : seuls les DOI/PMID présents dans les noms de fichiers "
            "seront utilisés.")
    os.makedirs(pdf_dir, exist_ok=True)
    ident = Identifier(http)
    rows = []          # (fichier, article|None, méthode, fiabilité, remarque, nouveau nom rel)
    seen = {}
    for i, path in enumerate(pdfs, 1):
        if STOP.is_set():
            log("⛔ Interruption : écriture des résultats partiels…")
            break
        try:
            info = PdfInfo(path)
            a, how, conf, note = ident.identify(info)
        except KeyboardInterrupt:
            log("⛔ Interruption : écriture des résultats partiels…")
            break
        rel = ""
        if a is not None:
            key = a.pmid or (a.doi or "").lower()
            if key and key in seen:
                note = f"doublon de « {seen[key]} » : non copié"
                conf = "doublon"
            elif conf == "sûre":
                seen[key] = info.name
                new = ident_filename(a)
                dest = os.path.join(pdf_dir, new)
                n = 2
                while os.path.exists(dest):
                    dest = os.path.join(pdf_dir, new[:-4] + f"_{n}.pdf")
                    n += 1
                shutil.copy2(path, dest)
                rel = f"{pdf_dirname}/{os.path.basename(dest)}"
            else:
                os.makedirs(check_dir, exist_ok=True)
                dest = os.path.join(check_dir, info.name)
                shutil.copy2(path, dest)
                rel = f"{pdf_dirname}/A_verifier/{info.name}"
        tag = (f"{a.pmid or 'NoPMID'} {a.first_author} ({a.year})" if a else "non identifié")
        mark = {"sûre": "✔", "à vérifier": "?", "doublon": "="}.get(conf, "✖")
        log(f"  [{i}/{len(pdfs)}] {mark} {info.name} → {tag}"
            + (f" — {how}" if how else "") + (f" ({note})" if note else ""))
        rows.append((info.name, a, how, conf, note, rel))

    sure = [r for r in rows if r[3] == "sûre"]
    refs = sorted((r[1] for r in sure), key=_sort_key)
    num = {id(a): i for i, a in enumerate(refs, 1)}

    # ---- Excel ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Références"
    hfill = PatternFill("solid", fgColor="1F4E78")
    headers = ["N° biblio", "Fichier d'origine", "PDF renommé (lien local)", "Identification",
               "Fiabilité", "PMID", "DOI", "PMCID", "Titre", "Journal", "Année", "Volume",
               "Numéro", "Pages", "Type", "1er auteur", "Dernier auteur", "Autres auteurs",
               "Référence (Vancouver)", "Remarque"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hfill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    order = sorted(rows, key=lambda r: (num.get(id(r[1]), 10 ** 6), r[0].lower()))
    for fname, a, how, conf, note, rel in order:
        if a is None:
            ws.append([clean_xl(v) for v in ["", fname, "non identifié", "", "", "", "", "", "", "",
                                             "", "", "", "", "", "", "", "", "", note]])
            ws.cell(ws.max_row, 3).font = Font(color="C00000")
            continue
        ws.append([clean_xl(v) for v in [
            num.get(id(a), ""), fname, rel or "-", how, conf, a.pmid or "NoPMID", a.doi, a.pmcid,
            a.title, a.journal, a.year, a.volume, a.issue, a.pages, "; ".join(a.pub_types),
            a.first_author, a.last_author, a.other_authors, vancouver(a), note]])
        r = ws.max_row
        if rel:
            ws.cell(r, 3).hyperlink = rel
            ws.cell(r, 3).font = Font(color="0563C1", underline="single")
        if a.pmid:
            ws.cell(r, 6).hyperlink = f"https://pubmed.ncbi.nlm.nih.gov/{a.pmid}/"
            ws.cell(r, 6).font = Font(color="0563C1", underline="single")
        if a.doi:
            ws.cell(r, 7).hyperlink = "https://doi.org/" + a.doi
            ws.cell(r, 7).font = Font(color="0563C1", underline="single")
        if conf != "sûre":
            ws.cell(r, 5).font = Font(bold=True, color="B26A00")
    for i, w in enumerate([9, 40, 44, 26, 11, 11, 26, 12, 60, 30, 7, 8, 8, 12, 22, 18, 18, 36, 90, 40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    n_sure, n_check = len(sure), sum(1 for r in rows if r[3] == "à vérifier")
    n_dup = sum(1 for r in rows if r[3] == "doublon")
    n_none = sum(1 for r in rows if r[1] is None)
    n_nopmid = sum(1 for r in sure if not r[1].pmid)
    ws2 = wb.create_sheet("Résumé")
    ws2.append(["Paramètre", "Valeur"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hfill
    meth = Counter(r[2] for r in sure)
    for k2, v in [("Dossier analysé", src_dir), ("Date", f"{now:%d/%m/%Y %H:%M}"),
                  ("PDF analysés", len(rows)), ("Identifiés (renommés)", n_sure),
                  ("  dont absents de PubMed (NoPMID)", n_nopmid),
                  ("À vérifier (copiés dans A_verifier, non renommés)", n_check),
                  ("Doublons (non copiés)", n_dup), ("Non identifiés", n_none),
                  ("Méthodes d'identification", ", ".join(f"{m2}: {c2}" for m2, c2 in meth.most_common())),
                  ("Style bibliographique", "Vancouver (NLM), ordre alphabétique du 1er auteur"),
                  ("Durée", fmt_duration(time.monotonic() - t0)),
                  ("Requêtes", f"{http.total} ({', '.join(f'{k3}: {v3}' for k3, v3 in http.counts.most_common())})"),
                  ("Script", f"pubmed_search.py v{VERSION}")]:
        ws2.append([clean_xl(k2), clean_xl(v)])
    ws2.column_dimensions["A"].width = 50
    ws2.column_dimensions["B"].width = 90
    xlsx = os.path.join(folder, f"{stem}.xlsx")
    wb.save(xlsx)

    # ---- Word ----
    docx = os.path.join(folder, f"{stem}.docx")
    write_docx(docx, "Bibliographie",
               f"{len(refs)} référence(s) — style Vancouver — {now:%d/%m/%Y} — dossier : {src_dir}",
               [vancouver(a) for a in refs])

    # ---- non identifiés / à vérifier ----
    todo = [r for r in rows if r[3] != "sûre"]
    if todo:
        with open(os.path.join(folder, "a_verifier.txt"), "w", encoding="utf-8") as f:
            f.write(f"PDF non renommés — {now:%d/%m/%Y %H:%M}\n"
                    "Pour un PDF non identifié : renommez l'original avec son PMID ou son DOI "
                    "(ex. 12345678.pdf ou 10.1000_xyz.pdf) puis relancez l'identification.\n\n")
            for fname, a, how, conf, note, rel in todo:
                prop = f" | proposition : {vancouver(a)}" if a is not None else ""
                f.write(f"[{conf or 'non identifié'}] {fname} — {note}{prop}\n")

    log("\n" + "=" * 60)
    log("📈 IDENTIFICATION")
    log("=" * 60)
    log(f"  PDF analysés            : {len(rows)}")
    log(f"  Identifiés et renommés  : {n_sure}" + (f"  (dont {n_nopmid} hors PubMed : NoPMID)" if n_nopmid else ""))
    log(f"  À vérifier              : {n_check}")
    log(f"  Doublons                : {n_dup}")
    log(f"  Non identifiés          : {n_none}")
    log(f"  Durée                   : {fmt_duration(time.monotonic() - t0)}")
    log(f"  Requêtes HTTP           : {http.total}")
    log(f"\n  Excel : {xlsx}")
    log(f"  Word  : {docx}")
    if todo:
        log(f"  Liste à vérifier : {os.path.join(folder, 'a_verifier.txt')}")
    return 0


# --------------------------------------------------------------------------- #
# Mode évaluation : combien d'articles, gratuits ou non, avant de lancer la recherche
# --------------------------------------------------------------------------- #
def _esearch_raw(http, term, sort="relevance", retmax=0, webenv=None, query_key=None):
    data = http.ncbi(db="pubmed", term=term, sort=sort, retmax=retmax, retmode="json")
    if webenv:
        data.update(WebEnv=webenv, usehistory="y")
    for wait in (5, 15, 30, 0):
        r = http.request("POST", EUTILS + "esearch.fcgi", "ncbi", data=data)
        r.raise_for_status()
        try:
            j = r.json().get("esearchresult", {})
        except ValueError:
            j = {"ERROR": "réponse illisible (backend NCBI)"}
        err = str(j.get("ERROR", "")).lower()
        if not (err and wait and any(k in err for k in NCBI_TRANSIENT)):
            break
        time.sleep(wait)
    if "ERROR" in j:
        raise RuntimeError(f"PubMed : {j['ERROR']}")
    return j


def evaluate_query(http, query, max_results, sorts):
    """Comptes PubMed d'une requête : tous les résultats et ceux qui seront traités
    (max_results premiers de chaque tri). « Gratuits » = filtre PubMed free full text[sb]
    (Free PMC article + Free article) ; « dont PMC » = "pubmed pmc"[sb] (aide PubMed)."""
    FREE, PMC = "free full text[sb]", "\"pubmed pmc\"[sb]"
    q = f"({query})"
    j = _esearch_raw(http, query)
    out = {"query": query, "total": int(j.get("count", 0)),
           "translation": j.get("querytranslation", ""), "warnings": []}
    for k in ("phrasesignored", "quotedphrasesnotfound", "outputmessages"):
        v = (j.get("warninglist") or {}).get(k)
        if v:
            out["warnings"].append(f"{k} : {', '.join(v) if isinstance(v, list) else v}")
    v = (j.get("errorlist") or {}).get("phrasesnotfound")
    if v:
        out["warnings"].append("termes non trouvés : " + ", ".join(v))
    out["free_total"] = int(_esearch_raw(http, f"{q} AND {FREE}").get("count", 0))
    out["pmc_total"] = int(_esearch_raw(http, f"{q} AND {PMC}").get("count", 0))
    ids = []
    for sort in sorts:
        ids += _esearch_raw(http, query, sort, max_results).get("idlist", [])
    ids = list(OrderedDict.fromkeys(ids))
    out["processed"] = len(ids)
    out["free_processed"] = out["pmc_processed"] = 0
    if ids:
        if out["processed"] == out["total"]:
            out["free_processed"], out["pmc_processed"] = out["free_total"], out["pmc_total"]
        else:
            r = http.request("POST", EUTILS + "epost.fcgi", "ncbi",
                             data=http.ncbi(db="pubmed", id=",".join(ids)))
            r.raise_for_status()
            root = ET.fromstring(r.content)
            env, key = root.findtext("WebEnv"), root.findtext("QueryKey")
            for name, sub in (("free_processed", FREE), ("pmc_processed", PMC)):
                out[name] = int(_esearch_raw(http, f"#{key} AND {sub}", webenv=env).get("count", 0))
    return out


def run_evaluate(args):
    http = Http(args.email, args.api_key, timeout=args.timeout)
    sorts = []
    if args.best_match or not args.most_recent:
        sorts.append("relevance")
    if args.most_recent:
        sorts.append("pub_date")
    try:
        ev = evaluate_query(http, args.query, args.max_results, sorts)
    except Exception as e:  # noqa
        print("EVAL_JSON: " + json.dumps({"query": args.query, "error": f"{type(e).__name__} : {e}"},
                                         ensure_ascii=False), flush=True)
        return 1
    ev["max"] = args.max_results
    ev["sorts"] = ["best" if x == "relevance" else "recent" for x in sorts]
    pct = lambda a, b: f"{100 * a / b:.0f} %" if b else "-"  # noqa: E731
    log(f"Évaluation de : {args.query}")
    log(f"  Résultats PubMed     : {ev['total']}  (gratuits {ev['free_total']} = {pct(ev['free_total'], ev['total'])}, "
        f"dont PMC {ev['pmc_total']} ; payants {ev['total'] - ev['free_total']})")
    log(f"  Articles à traiter   : {ev['processed']}  (gratuits {ev['free_processed']}, dont PMC "
        f"{ev['pmc_processed']} ; payants {ev['processed'] - ev['free_processed']})")
    for w in ev["warnings"]:
        log(f"  ⚠ {w}")
    print("EVAL_JSON: " + json.dumps(ev, ensure_ascii=False), flush=True)
    return 0


# --------------------------------------------------------------------------- #
# Programme principal
# --------------------------------------------------------------------------- #
SECRET_FLAGS = ("--api-key", "--openalex-key", "--elsevier-key", "--core-key", "--springer-oa-key",
                "--springer-meta-key", "--elsevier-insttoken")


def masked_argv():
    """Ligne de commande sans les clés API (pour error.txt)."""
    out, hide = [], False
    for x in sys.argv:
        if hide:
            out.append("***"); hide = False; continue
        if x in SECRET_FLAGS:
            hide = True
        elif x.startswith(tuple(f + "=" for f in SECRET_FLAGS)):
            x = x.split("=", 1)[0] + "=***"
        out.append(x)
    return " ".join(out)


# --------------------------------------------------------------------------- #
# Mode import : PDF récupérés à la main -> dossier de résultats
# --------------------------------------------------------------------------- #
def _norm(s):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


def _pdf_text(path, max_pages=3):
    if PdfReader is None:
        return ""
    try:
        r = PdfReader(path)
        return " ".join((pg.extract_text() or "") for pg in r.pages[:max_pages])
    except Exception:  # noqa
        return ""


def _name_from_excel(pmid, first_author, year):
    """'Fuchs LN' -> 22652600-Fuchs_LN(2025).pdf (même règle que la recherche)."""
    parts = (first_author or "").split()
    if len(parts) >= 2 and re.fullmatch(r"[A-Z]{1,4}", parts[-1]):
        who = f"{ascii_slug(' '.join(parts[:-1])) or 'Anonyme'}_{parts[-1]}"
    else:
        who = ascii_slug(first_author or "") or "Anonyme"
    return f"{pmid}-{who}({year or 'nd'}).pdf"


def run_import(src_dir, results_dir, move=False):
    src_dir = os.path.expanduser(src_dir)
    results_dir = os.path.expanduser(results_dir)
    xls = [x for x in glob.glob(os.path.join(results_dir, "*.xlsx")) if not os.path.basename(x).startswith("~$")]
    if len(xls) != 1:
        sys.exit(f"Dossier de résultats invalide : {len(xls)} fichier(s) .xlsx dans {results_dir}")
    xlsx = xls[0]
    stem = os.path.splitext(os.path.basename(xlsx))[0]
    pdf_dirname = f"PDF_{ascii_slug(stem, keep=r'A-Za-z0-9_-') or 'pubmed'}"
    pdf_dir = os.path.join(results_dir, pdf_dirname)
    os.makedirs(pdf_dir, exist_ok=True)
    if PdfReader is None:
        log("⚠ Module 'pypdf' absent (pip install pypdf) : reconnaissance limitée aux noms de "
            "fichiers contenant le PMID ou le DOI.")

    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    ws = wb["Articles"]
    hdr = {c.value: i for i, c in enumerate(ws[1], 1)}
    col = {k: hdr.get(k) for k in ("PMID", "DOI", "Titre de l'article", "1er auteur", "Année",
                                   "PDF (lien local)", "Source du PDF")}
    if not col["PMID"] or not col["PDF (lien local)"]:
        sys.exit("Onglet 'Articles' non reconnu.")
    missing = {}
    for r in range(2, ws.max_row + 1):
        link = str(ws.cell(r, col["PDF (lien local)"]).value or "")
        if link and link != "non téléchargé":
            continue
        g = lambda k: str(ws.cell(r, col[k]).value or "") if col[k] else ""  # noqa
        missing[g("PMID")] = {"row": r, "doi": g("DOI").lower(), "title": _norm(g("Titre de l'article")),
                              "author": g("1er auteur"), "year": g("Année")}
    log(f"Articles sans PDF dans {os.path.basename(xlsx)} : {len(missing)}")
    if not missing:
        return 0

    pdfs = sorted(glob.glob(os.path.join(src_dir, "*.pdf")) + glob.glob(os.path.join(src_dir, "*.PDF")))
    log(f"PDF trouvés dans {src_dir} : {len(pdfs)}")
    imported = []
    for path in pdfs:
        if not missing:
            break
        base = os.path.basename(path)
        base_l = base.lower()
        text = _pdf_text(path)
        text_l = text.lower()
        text_n = _norm(text)
        hit, how = None, ""
        for pmid, m in missing.items():
            if re.search(r"(?<!\d)%s(?!\d)" % pmid, base):
                hit, how = pmid, "PMID dans le nom du fichier"
                break
            if m["doi"] and (m["doi"] in text_l or m["doi"].replace("/", "_") in base_l
                             or (len(m["doi"].split("/")[-1]) >= 6 and m["doi"].split("/")[-1] in base_l)):
                hit, how = pmid, "DOI"
                break
            if len(m["title"]) >= 30 and m["title"][:80] in text_n:
                hit, how = pmid, "titre"
                break
        if not hit:
            continue
        m = missing.pop(hit)
        fname = _name_from_excel(hit, m["author"], m["year"])
        dest = os.path.join(pdf_dir, fname)
        (shutil.move if move else shutil.copy2)(path, dest)
        r = m["row"]
        cell = ws.cell(r, col["PDF (lien local)"])
        cell.value = f"{pdf_dirname}/{fname}"
        cell.hyperlink = f"{pdf_dirname}/{fname}"
        cell.font = Font(color="0563C1", underline="single")
        if col["Source du PDF"]:
            ws.cell(r, col["Source du PDF"]).value = "Import manuel"
        imported.append((hit, base, how))
        log(f"  ✔ {base} → {fname}  (reconnu par {how})")

    # mise à jour du résumé, des fichiers texte et de error.txt
    total = ws.max_row - 1
    n_pdf = sum(1 for r in range(2, ws.max_row + 1)
                if str(ws.cell(r, col["PDF (lien local)"]).value or "") not in ("", "non téléchargé"))
    ws2 = wb["Résumé"]
    for row in ws2.iter_rows(min_row=2):
        if row[0].value == "PDF téléchargés (total)":
            row[1].value = n_pdf
        elif row[0].value == "Taux de réussite / tous les articles":
            row[1].value = f"{100 * n_pdf / total:.1f} %" if total else "n/a"
    ws2.append([f"Import manuel du {dt.datetime.now():%d/%m/%Y %Hh%M}",
                f"{len(imported)} PDF ajouté(s) depuis {src_dir} : " + ", ".join(i[0] for i in imported)])
    wb.save(xlsx)
    sans = [str(ws.cell(r, col["PMID"]).value) for r in range(2, ws.max_row + 1)
            if str(ws.cell(r, col["PDF (lien local)"]).value or "") in ("", "non téléchargé")]
    for f in glob.glob(os.path.join(results_dir, "*_PMID_sans_pdf.txt")):
        with open(f, "w", encoding="utf-8") as fh:
            fh.write(",".join(sans))
    err = os.path.join(results_dir, "error.txt")
    if os.path.exists(err):
        with open(err, "a", encoding="utf-8") as fh:
            fh.write(f"\n{'-' * 78}\nIMPORT MANUEL ({dt.datetime.now():%d/%m/%Y %Hh%M}) : "
                     f"{len(imported)} PDF ajouté(s)\n")
            for pmid, base, how in imported:
                fh.write(f"  PMID {pmid} ← {base} (reconnu par {how})\n")
    log(f"\n{len(imported)} PDF importé(s) ; PDF au total : {n_pdf}/{total} ; "
        f"encore sans PDF : {len(sans)}")
    if missing:
        log("  Non retrouvés dans le dossier : " + ", ".join(missing))
    return 0


def check_ncbi_credentials(http):
    """Vérifie email et clé API NCBI ; sans clé valide, repasse à 3 requêtes/s."""
    if not EMAIL_RE.fullmatch(http.email or ""):
        log(f"⚠ Email '{http.email}' invalide : NCBI et Unpaywall exigent un email valide (--email).")
    if not http.api_key:
        log("ℹ Pas de clé API NCBI : débit limité à 3 requêtes/s (--api-key ou $NCBI_API_KEY).")
        return "non fournie (3 req/s)"
    try:
        r = http.request("GET", EUTILS + "einfo.fcgi", "ncbi", retries=1,
                         params=http.ncbi(db="pubmed", retmode="json"))
        body = r.text[:500].lower()
    except requests.RequestException as e:
        log(f"⚠ Vérification de la clé API NCBI impossible ({type(e).__name__}) ; clé conservée.")
        return "non vérifiée (10 req/s)"
    if r.status_code in (400, 401, 403) or "api key" in body and "invalid" in body:
        log("⚠ Clé API NCBI refusée par NCBI : poursuite SANS clé (3 requêtes/s).")
        http.api_key = None
        http.intervals["ncbi"] = 0.35
        return "refusée par NCBI → ignorée (3 req/s)"
    log("✔ Clé API NCBI valide : débit 10 requêtes/s.")
    return "valide (10 req/s)"


def parse_time_set(s):
    m = re.fullmatch(r"(\d{1,2})-(\d{1,2})-(\d{4})_(\d{1,2})[hH:](\d{2})", s.strip())
    if not m:
        raise argparse.ArgumentTypeError("format attendu : JJ-MM-AAAA_HHhMM (ex. 01-10-2026_20h45)")
    d, mo, y, h, mi = map(int, m.groups())
    try:
        return dt.datetime(y, mo, d, h, mi)
    except ValueError as e:
        raise argparse.ArgumentTypeError(str(e))


def wait_until(target):
    remaining = (target - dt.datetime.now()).total_seconds()
    if remaining <= 0:
        log(f"⚠ --time-set {target:%d-%m-%Y %Hh%M} est déjà passé : exécution immédiate.")
        return False
    log(f"⏳ Exécution programmée le {target:%d/%m/%Y à %Hh%M} "
        f"(dans {fmt_duration(remaining)}). Laisser la machine allumée ; Ctrl+C pour annuler.")
    while True:
        remaining = (target - dt.datetime.now()).total_seconds()
        if remaining <= 0:
            return True
        time.sleep(min(remaining, 30))


def parse_sources(s):
    items = [x.strip().lower() for x in (s or "").split(",") if x.strip()]
    if items in (["all"], ["toutes"]):
        return list(OPTIONAL_SOURCES)
    if items in (["none"], ["aucune"]):
        return []
    bad = [x for x in items if x not in OPTIONAL_SOURCES]
    if bad:
        raise argparse.ArgumentTypeError(f"source(s) inconnue(s) : {', '.join(bad)} ; "
                                         f"choisir parmi {', '.join(OPTIONAL_SOURCES)}")
    return items


def build_parser():
    p = argparse.ArgumentParser(
        description="Recherche PubMed et téléchargement des PDF d'articles en libre accès.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=('Exemple : python pubmed_search.py "diabetes[ti] AND review[pt]" '
                '--output Diabetes_review.xlsx --max-results 50 --getfreepaper'))
    p.add_argument("query", nargs="?", help='requête PubMed, ex. "diabetes[ti] AND review[pt]"')
    p.add_argument("--output", help="nom du fichier Excel, ex. Diabetes_review.xlsx (obligatoire "
                                    "pour une recherche)")
    p.add_argument("--max-results", type=int, default=100,
                   help="nombre maximal d'articles par tri (défaut 100, max 10000)")
    p.add_argument("--best-match", action="store_true", help="tri par pertinence (défaut)")
    p.add_argument("--most-recent", action="store_true",
                   help="tri par date de publication la plus récente")
    p.add_argument("--getfreepaper", action="store_true", help="télécharger les PDF")
    p.add_argument("--getjson", action="store_true",
                   help="passe JSON : texte intégral structuré (BioC PMC, Europe PMC, Springer Nature, "
                        "Elsevier) pour les articles restés sans PDF")
    p.add_argument("--getjats", action="store_true",
                   help="passe XML JATS pour tous les articles PMC : dossier JATS_<nom>/ pour une IA locale "
                        "(XML d'origine, Markdown, un CSV par tableau, figures ; suppléments en lien)")
    p.add_argument("--no-keywords", action="store_true",
                   help="ne pas créer <nom>_mots-cles.txt (mots-clés et MeSH regroupés en synonymes)")
    p.add_argument("--keywords-max-lookups", type=int, default=150,
                   help="mots-clés d'auteurs rattachés au MeSH par PubMed (1 requête chacun ; défaut 150)")
    p.add_argument("--keywords-from-queries", metavar="FICHIER",
                   help="mode liste : termes des requêtes de FICHIER (une par ligne) regroupés en synonymes "
                        "dans --output (défaut : FICHIER_mots-cles.txt)")
    p.add_argument("--sources", type=parse_sources, default=None,
                   help="sources optionnelles à utiliser, séparées par des virgules, parmi : "
                        + ", ".join(OPTIONAL_SOURCES) + " (défaut : toutes ; 'none' = aucune). "
                        "PMC S3, Europe PMC, HAL et sites des éditeurs restent toujours actifs.")
    p.add_argument("--pass-number", type=int, default=0,
                   help="nombre de passes supplémentaires sur les articles gratuits non téléchargés")
    p.add_argument("--time-set", type=parse_time_set,
                   help="lancer à la date/heure JJ-MM-AAAA_HHhMM (ex. 01-10-2026_20h45)")
    p.add_argument("--email", default=os.environ.get("NCBI_EMAIL", DEFAULT_EMAIL),
                   help="email transmis à NCBI (paramètre email des E-utilities) et à Unpaywall "
                        "(défaut : $NCBI_EMAIL, sinon sergesawadogo@gmail.com)")
    p.add_argument("--api-key", default=os.environ.get("NCBI_API_KEY"),
                   help="clé API NCBI, à créer dans 'Account settings' de son compte NCBI "
                        "(défaut : $NCBI_API_KEY) → 10 requêtes/s au lieu de 3")
    p.add_argument("--openalex-key", default=os.environ.get("OPENALEX_API_KEY"),
                   help="clé API OpenAlex (optionnelle, $OPENALEX_API_KEY)")
    p.add_argument("--elsevier-key", default=os.environ.get("ELSEVIER_API_KEY"),
                   help="clé API Elsevier gratuite (dev.elsevier.com, $ELSEVIER_API_KEY) : PDF des "
                        "articles libres ScienceDirect/Cell/JBC")
    p.add_argument("--elsevier-insttoken", default=os.environ.get("ELSEVIER_INSTTOKEN"),
                   help="jeton institutionnel Elsevier ($ELSEVIER_INSTTOKEN), fourni par Elsevier à une "
                        "institution abonnée : donne le texte intégral hors du réseau de l'institution")
    p.add_argument("--core-key", default=os.environ.get("CORE_API_KEY"),
                   help="clé API CORE gratuite (core.ac.uk/services/api, $CORE_API_KEY) : copies "
                        "des dépôts institutionnels")
    p.add_argument("--core-budget", type=float, default=5,
                   help="durée maximale (minutes) consacrée à CORE par exécution (défaut 5 ; "
                        "-1 = sans limite). Le quota gratuit de CORE ralentit fortement les recherches")
    p.add_argument("--springer-oa-key", default=os.environ.get("SPRINGER_OA_API_KEY"),
                   help="clé Springer Nature Open Access API (dev.springernature.com, "
                        "$SPRINGER_OA_API_KEY) : texte intégral JATS -> JSON")
    p.add_argument("--springer-meta-key", default=os.environ.get("SPRINGER_META_API_KEY"),
                   help="clé Springer Nature Meta API ($SPRINGER_META_API_KEY) : lien PDF des "
                        "articles Springer libres, résumé en dernier recours pour le JSON")
    p.add_argument("--import-pdf", metavar="DOSSIER",
                   help="mode import : range les PDF téléchargés à la main depuis DOSSIER "
                        "(ex. ~/Téléchargements) dans un dossier de résultats (--results)")
    p.add_argument("--results", metavar="DOSSIER_RESULTATS",
                   help="dossier de résultats à compléter avec --import-pdf "
                        "(ex. NETosis_pneumococcal_01.10.2026_18h46)")
    p.add_argument("--move", action="store_true",
                   help="avec --import-pdf : déplacer les PDF au lieu de les copier")
    p.add_argument("--outdir", default=None,
                   help="dossier parent des résultats (défaut : dossier courant ; avec "
                        "--identify-pdf : le dossier des PDF)")
    p.add_argument("--evaluate", action="store_true",
                   help="évaluer la requête sans rien télécharger : nombre de résultats PubMed, "
                        "gratuits (dont PMC) et payants, au total et parmi les articles qui seraient "
                        "traités (--max-results, tri). Dernière ligne : EVAL_JSON {…}")
    p.add_argument("--identify-pdf", metavar="DOSSIER",
                   help="identifier les PDF d'un dossier quelconque (DOI, PMID ou titre -> PubMed, "
                        "sinon Crossref) : copies renommées PMID-Auteur(Année).pdf, Excel des "
                        "références et bibliographie Word (Vancouver). Nom avec --output "
                        "(défaut Bibliographie)")
    p.add_argument("--workers", type=int, default=4, help="téléchargements parallèles (défaut 4)")
    p.add_argument("--timeout", type=int, default=40, help="délai max par requête en s (défaut 40)")
    return p


def main():
    argv = sys.argv[1:]
    if "--" in argv:  # un "--" isolé ferait prendre toutes les options suivantes pour la requête
        log("⚠ '--' isolé ignoré dans la ligne de commande.")
        argv = [x for x in argv if x != "--"]
    parser = build_parser()
    args = parser.parse_args(argv)
    if args.keywords_from_queries:
        return run_query_keywords(args)
    if args.import_pdf:
        if not args.results:
            parser.error("--import-pdf nécessite --results DOSSIER_RESULTATS")
        return run_import(args.import_pdf, args.results, move=args.move)
    if args.identify_pdf:
        http = Http(args.email, args.api_key, timeout=args.timeout)
        return run_identify(args.identify_pdf, args.output, args.outdir, http)
    args.outdir = args.outdir or "."
    if args.evaluate:
        if not args.query:
            parser.error("--evaluate nécessite la requête")
        if not 1 <= args.max_results <= 10000:
            sys.exit("--max-results doit être compris entre 1 et 10000 (limite ESearch de NCBI).")
        return run_evaluate(args)
    if not args.query or not args.output:
        parser.error("une recherche nécessite la requête et --output")
    if not args.output.lower().endswith(".xlsx"):
        args.output += ".xlsx"
    stem = ascii_slug(os.path.splitext(os.path.basename(args.output))[0], keep=r"A-Za-z0-9_\-") or "pubmed"
    if not 1 <= args.max_results <= 10000:
        sys.exit("--max-results doit être compris entre 1 et 10000 (limite ESearch de NCBI).")
    if args.pass_number < 0:
        sys.exit("--pass-number doit être ≥ 0.")
    if args.pass_number and not args.getfreepaper:
        log("⚠ --pass-number ignoré sans --getfreepaper.")
    if args.time_set:
        waited = wait_until(args.time_set)
        run_dt = args.time_set if waited else dt.datetime.now()
    else:
        run_dt = dt.datetime.now()

    folder_name = f"{stem}_{run_dt:%d.%m.%Y_%Hh%M}"
    folder = os.path.join(args.outdir, folder_name)
    k = 2
    while os.path.exists(folder):
        folder = os.path.join(args.outdir, f"{folder_name}_{k}")
        k += 1
    pdf_dirname = f"PDF_{stem}"
    pdf_dir = os.path.join(folder, pdf_dirname)
    os.makedirs(pdf_dir if args.getfreepaper else folder, exist_ok=True)
    json_dirname = f"JSON_{stem}"
    json_dir = os.path.join(folder, json_dirname)
    if args.getjson:
        os.makedirs(json_dir, exist_ok=True)
    jats_dirname = f"JATS_{stem}"
    jats_dir = os.path.join(folder, jats_dirname)
    if args.getjats:
        os.makedirs(jats_dir, exist_ok=True)
        with open(os.path.join(jats_dir, "LISEZMOI.txt"), "w", encoding="utf-8") as f:
            f.write(JATS_README)

    http = Http(args.email, args.api_key, args.openalex_key, args.timeout, args.elsevier_key,
                args.core_key, args.springer_oa_key, args.springer_meta_key, args.sources)
    http.elsevier_insttoken = args.elsevier_insttoken if args.elsevier_key else None
    http.core_budget = -1 if args.core_budget < 0 else int(args.core_budget * 60)
    t0 = time.monotonic()
    log(f"pubmed_search.py v{VERSION} — {run_dt:%d/%m/%Y %Hh%M}")
    log(f"Requête : {args.query}")
    log(f"Dossier : {os.path.abspath(folder)}")
    log(f"Email NCBI/Unpaywall : {args.email}")
    key_status = check_ncbi_credentials(http)
    off = [x for x in OPTIONAL_SOURCES if not http.use(x)]
    if off:
        log("ℹ Sources désactivées : " + ", ".join(off))

    sorts = []
    if args.best_match or not args.most_recent:
        sorts.append(("best", "relevance", "Best match (pertinence)"))
    if args.most_recent:
        sorts.append(("recent", "pub_date", "Most recent (date de publication)"))

    ranks, total_count, translation = {}, 0, ""
    for key, sort, label in sorts:
        log(f"\n🔎 ESearch — tri {label}, max {args.max_results}")
        count, ids, qt = esearch(http, args.query, sort, args.max_results)
        total_count, translation = count, qt or translation
        ranks[key] = ids
        log(f"  {count} résultats dans PubMed, {len(ids)} récupérés")

    order = list(OrderedDict.fromkeys(ranks.get("best", []) + ranks.get("recent", [])))
    best_set, recent_set = set(ranks.get("best", [])), set(ranks.get("recent", []))
    overlap = best_set & recent_set
    if len(sorts) == 2:
        log("\n📊 Comparaison des tris :")
        log(f"  uniques pertinence (best match seulement) : {len(best_set - recent_set)}")
        log(f"  uniques plus récents (most recent seulement) : {len(recent_set - best_set)}")
        log(f"  chevauchants (dans les deux) : {len(overlap)}")
        log(f"  total d'articles distincts traités : {len(order)}")

    articles = []
    if order:
        log("\n📥 Récupération des métadonnées (EFetch)…")
        recs = efetch_articles(http, order)
        for pmid in order:
            a = recs.get(pmid) or Article(pmid)
            articles.append(a)
        by_pmid = {a.pmid: a for a in articles}
        for i, p in enumerate(ranks.get("best", []), 1):
            by_pmid[p].rank_best = i
        for i, p in enumerate(ranks.get("recent", []), 1):
            by_pmid[p].rank_recent = i
        log("🔓 Détection des articles gratuits (filtre PubMed 'free full text')…")
        free = free_fulltext_set(http, order)
        for a in articles:
            if a.pmid in free:
                a.free = "Free PMC article" if a.pmcid else "Free article"
            elif a.pmcid:
                a.free = "PMC (sous embargo ?)"

    passes_done = 0
    dl = None
    jf = None
    jx = None
    interrupted = False
    if (args.getfreepaper or args.getjson or args.getjats) and articles:
        try:
            log("🔗 Enrichissement des identifiants et des liens (ID converter PMC, Europe PMC, "
                "LinkOut" + (", OpenAlex" if args.openalex_key else "") + ")…")
            if args.core_key:
                log("ℹ CORE activé : ~10 requêtes/min (limite de l'API), seulement pour les articles "
                    "non trouvés par les autres sources"
                    + (f", {args.core_budget:g} min au plus." if args.core_budget >= 0 else "."))
            log("  • ID converter PMC (PMCID, DOI)…")
            idconv_fill(http, articles)
            log("  • Europe PMC (liens libres)…")
            europepmc_enrich(http, articles)
            if args.getfreepaper:
                log("  • Liens éditeurs PubMed (LinkOut)…")
                elink_prlinks(http, [a for a in articles if a.free != "Non" or not a.pmcid])
                if args.openalex_key:
                    log("  • OpenAlex…")
                openalex_enrich(http, articles)
            # le PMCID a pu être découvert après coup
            for a in articles:
                if a.pmcid and a.free == "Free article":
                    a.free = "Free PMC article"
            dl = Downloader(http, pdf_dir, pdf_dirname)
            for p in (range(1, args.pass_number + 2) if args.getfreepaper else []):
                if p == 1:
                    targets = articles
                else:
                    targets = [a for a in articles if not a.pdf_rel and a.free_like]
                    if not targets:
                        break
                    pause = min(20 * (p - 1), 90)
                    log(f"\n⏸ Pause de {pause}s avant la passe {p}…")
                    time.sleep(pause)
                log(f"\n⬇ Passe {p} : {len(targets)} article(s) à traiter")
                passes_done = p
                n = len(targets)
                with ThreadPoolExecutor(max_workers=max(1, args.workers)) as ex:
                    futs = {ex.submit(dl.process, a, p): a for a in targets}
                    for i, fut in enumerate(as_completed(futs), 1):
                        a = futs[fut]
                        try:
                            fut.result()
                        except Exception as e:  # noqa
                            a.note(p, "script", "", f"exception : {type(e).__name__}: {e}")
                        status = f"OK ({a.pdf_source})" if a.pdf_pass == p else \
                            ("déjà téléchargé" if a.pdf_rel else "non téléchargé")
                        log(f"  [{i}/{n}] PMID {a.pmid} [{a.free}] → {status}")
                ok_now = sum(1 for a in articles if a.pdf_rel)
                log(f"  Fin de passe {p} : {ok_now} PDF au total")
            if args.getjson and not STOP.is_set():
                jf = JsonFetcher(http, json_dir, json_dirname, dl)
                targets = [a for a in articles if not a.pdf_rel]
                log(f"\n🧾 Passe JSON : {len(targets)} article(s) sans PDF")
                n = len(targets)
                with ThreadPoolExecutor(max_workers=max(1, args.workers)) as ex:
                    futs = {ex.submit(jf.process, a): a for a in targets}
                    for i, fut in enumerate(as_completed(futs), 1):
                        a = futs[fut]
                        try:
                            fut.result()
                        except Exception as e:  # noqa
                            a.jnote("script", "", f"exception : {type(e).__name__}: {e}")
                        status = f"JSON {a.json_level} ({a.json_source})" if a.json_rel else "pas de JSON"
                        log(f"  [{i}/{n}] PMID {a.pmid} [{a.free}] ⇒ {status}")
                log(f"  Fin de la passe JSON : {sum(1 for a in articles if a.json_rel)} JSON")
            if args.getjats and not STOP.is_set():
                jx = JatsFetcher(http, jats_dir, jats_dirname)
                targets = [a for a in articles if a.pmcid]
                log(f"\n🧬 Passe XML JATS (IA locale) : {len(targets)} article(s) PMC "
                    f"({sum(1 for a in articles if not a.pmcid)} sans PMCID, pas de XML libre)")
                n = len(targets)
                with ThreadPoolExecutor(max_workers=max(1, args.workers)) as ex:
                    futs = {ex.submit(jx.process, a): a for a in targets}
                    for i, fut in enumerate(as_completed(futs), 1):
                        a = futs[fut]
                        try:
                            fut.result()
                        except Exception as e:  # noqa
                            a.jats_note = f"exception : {type(e).__name__}: {e}"
                        status = f"XML ({a.jats_source}) : {a.jats_info}" if a.jats_rel else \
                            f"pas de XML ({(a.jats_note or '')[:90]})"
                        log(f"  [{i}/{n}] PMID {a.pmid} ◆ {status}")
                log(f"  Fin de la passe XML JATS : {sum(1 for a in articles if a.jats_rel)} article(s), "
                    f"{jx.n_fig} figure(s), {jx.n_tab} tableau(x) CSV")
        except KeyboardInterrupt:
            STOP.set()
            interrupted = True
            log("\n⛔ Interruption : écriture des résultats partiels…")

    kw_groups = None
    kw_path = os.path.join(folder, f"{stem}_mots-cles.txt")
    if articles and not args.no_keywords:
        log("\n🏷 Mots-clés et termes MeSH équivalents" + (" (sans requête réseau : arrêt demandé)"
                                                          if interrupted or STOP.is_set() else "") + "…")
        try:
            kw_groups = build_keywords(http, articles, kw_path, args.query, args.keywords_max_lookups,
                                       offline=interrupted or STOP.is_set())
        except KeyboardInterrupt:
            interrupted = True
        except Exception as e:  # noqa
            log(f"  ⚠ Mots-clés : {type(e).__name__} : {e}")

    # ---------------- statistiques et sorties ---------------- #
    elapsed = time.monotonic() - t0
    n = len(articles)
    n_pdf = sum(1 for a in articles if a.pdf_rel)
    free_cnt = Counter(a.free for a in articles)
    n_free = sum(1 for a in articles if a.free in ("Free PMC article", "Free article"))
    n_free_pdf = sum(1 for a in articles if a.pdf_rel and a.free in ("Free PMC article", "Free article"))
    n_freelike = sum(1 for a in articles if a.free_like)
    n_freelike_pdf = sum(1 for a in articles if a.free_like and a.pdf_rel)
    remaining = sum(1 for a in articles if a.free_like and not a.pdf_rel)
    pct = lambda x, y: f"{100 * x / y:.1f} %" if y else "n/a"  # noqa

    with open(os.path.join(folder, f"{stem}_PMID_tous.txt"), "w", encoding="utf-8") as f:
        f.write(",".join(a.pmid for a in articles))
    with open(os.path.join(folder, f"{stem}_PMID_sans_pdf.txt"), "w", encoding="utf-8") as f:
        f.write(",".join(a.pmid for a in articles if not a.pdf_rel))
    n_json = sum(1 for a in articles if a.json_rel)
    n_json_ft = sum(1 for a in articles if a.json_level == "texte intégral")
    n_jats = sum(1 for a in articles if a.jats_rel)
    n_pmc = sum(1 for a in articles if a.pmcid)
    if args.getjson:
        with open(os.path.join(folder, f"{stem}_PMID_json.txt"), "w", encoding="utf-8") as f:
            f.write(",".join(a.pmid for a in articles if a.json_rel))

    req_detail = ", ".join(f"{k}: {v}" for k, v in http.counts.most_common())
    src_detail = ", ".join(f"{k}: {v}" for k, v in (dl.source_ok.most_common() if dl else []))
    summary = [
        ("Requête PubMed", args.query),
        ("Traduction PubMed de la requête", translation),
        ("Date/heure d'exécution", f"{run_dt:%d/%m/%Y %Hh%M}"),
        ("Tri(s)", " + ".join(s[2] for s in sorts)),
        ("--max-results (par tri)", args.max_results),
        ("Nombre total de résultats dans PubMed", total_count),
        ("Articles distincts récupérés", n),
    ]
    if len(sorts) == 2:
        summary += [("Uniques pertinence (best match seulement)", len(best_set - recent_set)),
                    ("Uniques plus récents (most recent seulement)", len(recent_set - best_set)),
                    ("Chevauchants (les deux tris)", len(overlap))]
    summary += [
        ("Free PMC article", free_cnt.get("Free PMC article", 0)),
        ("Free article", free_cnt.get("Free article", 0)),
        ("PMC sous embargo ?", free_cnt.get("PMC (sous embargo ?)", 0)),
        ("Sans mention de gratuité", free_cnt.get("Non", 0)),
        ("Articles gratuits ou avec version OA détectée", n_freelike if args.getfreepaper else n_free),
        ("Téléchargement demandé (--getfreepaper)", "oui" if args.getfreepaper else "non"),
        ("PDF téléchargés (total)", n_pdf),
        ("  dont articles 'Free' PubMed", n_free_pdf),
        ("Taux de réussite / articles 'Free' PubMed", pct(n_free_pdf, n_free)),
        ("Taux de réussite / gratuits + OA détectés", pct(n_freelike_pdf, n_freelike)),
        ("Taux de réussite / tous les articles", pct(n_pdf, n)),
        ("Gratuits/OA restant à télécharger", remaining),
        ("PDF par source", src_detail),
        ("Passes effectuées", passes_done),
        ("Passe JSON demandée (--getjson)", "oui" if args.getjson else "non"),
        ("JSON obtenus (articles sans PDF)", n_json if args.getjson else "-"),
        ("  dont texte intégral", n_json_ft if args.getjson else "-"),
        ("  dont résumé seulement", (n_json - n_json_ft) if args.getjson else "-"),
        ("JSON par source", ", ".join(f"{k}: {v}" for k, v in jf.source_ok.most_common()) if jf else "-"),
        ("PDF ou JSON texte intégral / tous les articles", pct(n_pdf + n_json_ft, n)),
        ("Passe XML JATS demandée (--getjats)", "oui" if args.getjats else "non"),
        ("XML JATS obtenus / articles PMC", f"{n_jats} / {n_pmc}" if args.getjats else "-"),
        ("XML JATS par source", ", ".join(f"{k}: {v}" for k, v in jx.source_ok.most_common()) if jx else "-"),
        ("Figures et tableaux CSV (IA locale)", f"{jx.n_fig} figure(s), {jx.n_tab} tableau(x), "
                                                f"{jx.n_tab_img} tableau(x) en image" if jx else "-"),
        ("Lignes de mots-clés équivalents", len(kw_groups) if kw_groups else "-"),
        ("Sources optionnelles actives", ", ".join(x for x in OPTIONAL_SOURCES if http.use(x)) or "aucune"),
        ("Durée", fmt_duration(elapsed)),
        ("Requêtes HTTP (total)", http.total),
        ("Requêtes HTTP par service", req_detail),
        ("Email transmis à NCBI/Unpaywall", args.email),
        ("Clé API NCBI", key_status),
        ("Clé API OpenAlex", "fournie" if args.openalex_key else "non fournie (OpenAlex non interrogé)"),
        ("Clé API Elsevier", ("fournie" + (" + jeton institutionnel" if http.elsevier_insttoken else "")
                              + (" — refusée pour le texte intégral (pas de droits)" if http.els_denied else ""))
                             if args.elsevier_key else "non fournie (API Elsevier non interrogée)"),
        ("Clé API CORE", "fournie" if args.core_key else "non fournie (CORE non interrogé)"),
        ("Clé Springer Nature Open Access", "fournie" if args.springer_oa_key else "non fournie"),
        ("Clé Springer Nature Meta", "fournie" if args.springer_meta_key else "non fournie"),
        ("Interrompu", "oui" if interrupted else "non"),
    ]
    xlsx_path = os.path.join(folder, os.path.basename(args.output))
    write_excel(xlsx_path, articles, summary, kw_groups)
    meta = [f"Date : {run_dt:%d/%m/%Y %Hh%M}", f"Requête : {args.query}",
            f"Commande : {masked_argv()}", f"Version script : {VERSION} | Python {sys.version.split()[0]}"
            f" | requests {requests.__version__}",
            f"Passes : {passes_done} | PDF : {n_pdf}/{n} | Requêtes : {http.total} ({req_detail})"]
    tips = advice(articles, http, args.getfreepaper)
    write_error_file(os.path.join(folder, "error.txt"), articles, meta, args.getfreepaper or args.getjson, tips,
                     getjats=args.getjats)

    log("\n" + "=" * 60)
    log("📈 STATISTIQUES")
    log("=" * 60)
    log(f"  Résultats PubMed (total)       : {total_count}")
    log(f"  Articles traités               : {n}")
    log(f"  Gratuits (Free PMC / Free)     : {free_cnt.get('Free PMC article', 0)} / "
        f"{free_cnt.get('Free article', 0)}")
    if args.getfreepaper:
        log(f"  Gratuits + OA détectés         : {n_freelike}")
        extra = n_pdf - n_freelike_pdf
        log(f"  PDF téléchargés                : {n_pdf}  ({pct(n_freelike_pdf, n_freelike)} des "
            f"gratuits/OA" + (f" ; dont {extra} article(s) non signalé(s) gratuit(s)" if extra else "")
            + ")")
        log(f"  Restant à télécharger          : {remaining}")
        log(f"  PDF par source                 : {src_detail or '-'}")
        log(f"  Passes effectuées              : {passes_done}")
    if args.getjson:
        log(f"  JSON obtenus                   : {n_json}  (texte intégral : {n_json_ft}, "
            f"résumé : {n_json - n_json_ft})")
        log(f"  JSON par source                : "
            f"{', '.join(f'{k}: {v}' for k, v in jf.source_ok.most_common()) if jf else '-'}")
    if args.getjats:
        log(f"  XML JATS (IA locale)           : {n_jats} / {n_pmc} articles PMC"
            + (f"  ({jx.n_fig} figures, {jx.n_tab} tableaux CSV)" if jx else ""))
    if kw_groups:
        log(f"  Mots-clés équivalents          : {len(kw_groups)} ligne(s) → {stem}_mots-cles.txt")
    for t in tips:
        log(f"  💡 {t}")
    log(f"  Durée                          : {fmt_duration(elapsed)}")
    log(f"  Requêtes HTTP                  : {http.total}  ({req_detail})")
    if http.paused_log:
        log("  Services mis en pause          : " + ", ".join(
            f"{BREAKER_SERVICES.get(k, k)} ({http.skipped.get(k, 0)} requête(s) évitée(s))"
            for k in sorted({k for k, _ in http.paused_log})))
    log(f"\n  Excel : {xlsx_path}")
    log(f"  Fichiers : {stem}_PMID_tous.txt, {stem}_PMID_sans_pdf.txt"
        + (f", {stem}_PMID_json.txt" if args.getjson else "")
        + (f", {stem}_mots-cles.txt" if kw_groups else "") + ", error.txt"
        + (f", {jats_dirname}/" if args.getjats else ""))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        log("\n⛔ Annulé.")
        sys.exit(130)
    except requests.RequestException as e:
        log(f"\n❌ Erreur réseau : {e}")
        sys.exit(2)
    except RuntimeError as e:
        log(f"\n❌ {e}")
        sys.exit(2)
